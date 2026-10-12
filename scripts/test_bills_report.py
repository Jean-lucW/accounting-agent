#!/usr/bin/env python
"""Offline checks for the bills payable report: every status and every
classifier rule, on fake bills and fake bank lines.

No network, no credentials, no Xero: exchange rates are a fixed table, the
bank and the reconcile list are fixtures. Runs against config/group.example.toml.

    .venv/bin/python scripts/test_bills_report.py
"""

from __future__ import annotations

import json
import os
import re
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
os.environ["AGENT_GROUP_CONFIG"] = str(ROOT / "config" / "group.example.toml")

import bills_report as br  # noqa: E402
from accounting_agent import config  # noqa: E402

FAILS: list[str] = []


def check(cond: bool, what: str) -> None:
    print(("ok    " if cond else "FAIL  ") + what)
    if not cond:
        FAILS.append(what)


# A fixed rate table instead of the ECB: no network.
RATES = {("EUR", "GBP"): 0.86, ("GBP", "EUR"): 1 / 0.86, ("EUR", "USD"): 1.1, ("GBP", "USD"): 1.3}
br._fx = lambda base, quote, day: 1.0 if base == quote else RATES.get((base, quote))

KEY = "opco-us"
OTHER = "holdco"
_n = [0]


def bill(supplier, day, amount, ccy="USD", number="", desc="", status="OPEN", due=None, **kw):
    _n[0] += 1
    b = {"id": f"id{_n[0]}", "status": status, "date": day, "due": due or day, "supplier": supplier,
         "number": number or f"INV-{_n[0]}", "reference": "", "currency": ccy, "total": amount, "paid": 0.0,
         "credited": 0.0, "due_amount": amount, "accounts": [], "account_codes": [], "description": desc,
         "attachment": True, "updated": day, "fully_paid_on": "", "payments": []}
    b.update(kw)
    return b


def line(day, payee, amount, ccy="USD", entity=KEY, reference="", who="", orig=None):
    return {"entity": entity, "account": "Main", "date": day, "payee": payee, "amount": -abs(amount),
            "currency": ccy, "orig_amount": orig[1] if orig else None, "orig_currency": orig[0] if orig else "",
            "who": who, "reference": reference, "source": "fixture"}


def spend(day, amount, accounts=("Office costs",)):
    return {"date": day, "amount": amount, "currency": "USD", "bank_account": "Main", "contact": "",
            "reference": "", "accounts": list(accounts), "reconciled": False}


def run(b, history=(), others=(), spends=(), paid=(), batch=None, claimed=None):
    ent = {"bills": [b, *others, *paid], "spends": list(spends)}
    store = {"entities": {KEY: ent}}
    return br.classify(KEY, b, list(history), [], store, {}, batch, claimed if claimed is not None else set())


# ---- the base case ------------------------------------------------------------

st, why, _ = run(bill("Example Supplier Ltd", "2026-06-10", 480.0))
check(st == "UNPAID", "no bank line: UNPAID")

st, why, ev = run(bill("Example Supplier Ltd", "2026-06-10", 480.0),
                  [line("2026-06-10", "Example Supplier", 480.0)])
check(st == "MATCH IN XERO" and "Example Supplier" in ev, "own bank, exact amount, name: MATCH IN XERO")

# ---- NOT APPROVED -------------------------------------------------------------

approved = bill("Example Supplier Ltd", "2026-06-10", 480.0, number="INV-7001")
copy_no = bill("Example Supplier Ltd", "2026-06-20", 99.0, number="inv 7001", status="DRAFT")
copy_amt = bill("Example Supplier Ltd.", "2026-06-14", 480.0, number="X-1", status="SUBMITTED")
plain = bill("Another Vendor", "2026-06-14", 55.0, status="DRAFT")
pool = [approved, copy_no, copy_amt, plain]
why_no = br._draft_why(copy_no, pool)
check("do not pay" in why_no and "void" in why_no, "draft with an approved bill's number: do not pay, void")
why_amt = br._draft_why(copy_amt, pool)
check("do not pay" in why_amt and "submitted" in why_amt, "draft, same supplier, currency, amount within 7 days: do not pay")
far = bill("Example Supplier Ltd", "2026-07-01", 480.0, status="DRAFT")
check("approve it before paying" in br._draft_why(far, [approved, far]), "same amount 21 days apart is not a copy")
check("approve it before paying" in br._draft_why(plain, pool), "an unrelated draft: approve it before paying")

# assess() puts them in the status and the counts, without touching the bank
real_history, real_feed = br.bank_history, br.bank_feed_lines
br.bank_history = lambda: ([], {})
br.bank_feed_lines = lambda: []
store = {"entities": {KEY: {"bills": [approved, copy_no, plain], "spends": []}}}
br.assess(store)
check(copy_no["state"] == "NOT APPROVED" and plain["state"] == "NOT APPROVED", "assess marks drafts NOT APPROVED")
check(store["counts"][KEY]["NOT APPROVED"] == 2 and store["counts"][KEY]["UNPAID"] == 1, "counts carry NOT APPROVED")
check("NOT APPROVED" in br.STATE_ORDER and "NOT APPROVED" in br.STATE_HELP, "status order and legend carry NOT APPROVED")

# ---- MONTHLY DEBIT ------------------------------------------------------------

br.PERIODIC_BILLING = [(KEY, re.compile("example rides", re.I), "2026-09-01", "on the monthly collection")]
st, why, _ = run(bill("Example Rides", "2026-09-05", 40.0))
check(st == "MONTHLY DEBIT" and "monthly collection" in why, "supplier on a monthly debit from its date: MONTHLY DEBIT")
st, _, _ = run(bill("Example Rides", "2026-08-30", 40.0))
check(st == "UNPAID", "a bill before the from date is classified normally")
st, _, _ = run(bill("Example Rides", "2026-09-05", 40.0), [line("2026-09-05", "Example Rides", 40.0)])
check(st == "MONTHLY DEBIT", "a monthly debit bill never reads as paid early by a card line")
br.PERIODIC_BILLING = []

# ---- never_duplicate ----------------------------------------------------------

order_a = bill("Example Takeaway", "2026-06-10", 25.0, number="A-1")
order_b = bill("Example Takeaway", "2026-06-11", 25.0, number="A-2", status="PAID",
               payments=[{"date": "2026-06-11", "amount": 25.0, "account": "Main", "account_code": "090",
                          "is_bank": True, "reconciled": True, "reference": ""}])
hist = [line("2026-06-11", "Example Takeaway", 25.0)]
st, _, _ = run(order_a, hist, paid=[order_b])
check(st == "DUPLICATE?", "a line already paid against another bill is DUPLICATE? by default")
br.NEVER_DUPLICATE = re.compile("example takeaway", re.I)
st, why, _ = run(order_a, hist, paid=[order_b])
check(st == "UNPAID" and "no duplicate" in why, "never_duplicate supplier: repeat order at one price is UNPAID")
br.NEVER_DUPLICATE = None

# ---- a named, exact-amount, unclaimed line may precede the bill (rent ahead) ----

rent = lambda amt=2000.0: bill("Example Landlord Ltd", "2026-07-01", amt, desc="Rent July")   # noqa: E731
ahead = line("2026-06-08", "Example Landlord", 2000.0)
st, _, _ = run(rent(), [ahead])
check(st == "MATCH IN XERO", "named exact line 23 days before a bill of 1,000 or more: matched")
st, _, _ = run(rent(), [line("2026-05-15", "Example Landlord", 2000.0)])
check(st == "UNPAID", "beyond named_lead_days the line is not the payment")
st, _, _ = run(rent(500.0), [line("2026-06-08", "Example Landlord", 500.0)])
check(st == "UNPAID", "below named_lead_min only lead_days applies")
st, _, _ = run(rent(), [ahead], spends=[spend("2026-06-08", 2000.0)])
check(st != "MATCH IN XERO", "a line Xero already holds is not taken as payment ahead")
st, _, _ = run(rent(), [line("2026-06-08", "Stranger Cafe", 2000.0)])
check(st == "UNPAID", "an early line must name the supplier")

# ---- named_amount_tolerance: a reimbursed claim at a slightly different rate ----

claim = lambda: bill("Example Supplier Ltd", "2026-06-10", 1000.0, ccy="EUR")   # noqa: E731
st, _, _ = run(claim(), [line("2026-06-10", "To Example Supplier", 1004.0, ccy="EUR")])
check(st == "MATCH IN XERO", "named transfer within 1% in the bill's currency: the bill")
st, _, _ = run(claim(), [line("2026-06-10", "To Example Supplier", 1030.0, ccy="EUR")])
check(st == "UNPAID", "named transfer 3% out: not the bill")

# ---- agent payees ---------------------------------------------------------------

travel = lambda: bill("Acme Travel Ltd", "2026-06-10", 300.0, desc="Accommodation")   # noqa: E731
hotel_line = line("2026-06-10", "Hotelsite Booking", 300.0)
st, why, _ = run(travel(), [hotel_line])
check(st == "MATCH IN XERO" and "confirm" in why, "no agent table: payee unlike the supplier needs confirming")
br.AGENT_PAYEES = {"acme": {"hotelsite"}}
check(br._bill_score(hotel_line, travel()) == 2, "agent payee table: the merchant reads as the bill")
st, why, _ = run(travel(), [hotel_line])
check(st == "MATCH IN XERO" and "confirm" not in why, "agent payee table: matched without a question")
br.AGENT_PAYEES = {}

# ---- a person's transfer matched on what the bill is for ------------------------

rent_bill = lambda: bill("Example Property Ltd", "2026-06-10", 1500.0, desc="Office rent October")   # noqa: E731
reimb = line("2026-06-10", "To Sam Taylor", 1500.0, reference="October office rent")
check(br._bill_score(reimb, rent_bill()) == 1, "transfer to a person whose reference overlaps the description scores 1")
claimed: set = set()
st, _, _ = run(rent_bill(), [reimb], claimed=claimed)
check(st == "MATCH IN XERO" and not claimed, "reimbursed transfer matches and is never claimed")
claimed2: set = set()
st, _, _ = run(bill("Example Supplier Ltd", "2026-06-10", 480.0), [line("2026-06-10", "Example Supplier", 480.0)], claimed=claimed2)
check(len(claimed2) == 1, "an ordinary match is claimed")

# ---- same_charge: a card charge under another merchant's name -------------------

op = lambda day="2026-06-10": bill("Example Operator Ltd", day, 100.0, ccy="EUR", desc="Team dinner")   # noqa: E731
st, why, _ = run(op(), [line("2026-06-11", "Mystery Merchant", 86.3, ccy="GBP")])
check(st == "MATCH IN XERO" and "same card" in why, "stranger payee, within 2%, one day after: the bill")
st, why, _ = run(op(), [line("2026-06-11", "Mystery Merchant", 88.4, ccy="GBP")])
check(st == "UNPAID" and "another payee" in why, "stranger payee 2.8% out at the rate: UNPAID")
st, _, _ = run(op(), [line("2026-06-15", "Mystery Merchant", 86.3, ccy="GBP")])
check(st == "UNPAID", "stranger payee beyond stranger_days: UNPAID")
st, _, _ = run(op(), [line("2026-06-11", "Mystery Merchant", 86.3, ccy="GBP", who="Pat Someone")])
check(st == "UNPAID", "another person's card than the bill names: UNPAID")
nearer = bill("Other Vendor", "2026-06-10", 100.5, ccy="EUR")
st, _, _ = run(op(), [line("2026-06-11", "Mystery Merchant", 86.4, ccy="GBP")], others=[nearer])
check(st == "UNPAID", "another open bill fits the line better: UNPAID")
st, _, _ = run(op(), [line("2026-06-11", "Mystery Merchant", 86.3, ccy="GBP")], spends=[spend("2026-06-11", 86.3)])
check(st == "UNPAID", "a line Xero already holds is not a stranger's charge")

# ---- batch payments ------------------------------------------------------------

b1 = bill("Example Supplier Ltd", "2026-06-10", 100.0, number="B-1")
b2 = bill("Example Supplier Ltd", "2026-06-12", 150.0, number="B-2")
open_by = {KEY: [b1, b2]}
got = br._batches(open_by, [line("2026-06-15", "Example Supplier", 250.0)])
check(set(got) == {b1["id"], b2["id"]}, "one named payment after the bills covers both")
check(not br._batches(open_by, [line("2026-05-20", "Example Supplier", 250.0)]),
      "a payment weeks before the earliest bill is not their batch")
check(not br._batches(open_by, [line("2026-06-15", "Stranger Cafe", 250.0)]),
      "a batch payment needs the supplier's name")
st, why, _ = run(b1, [line("2026-06-15", "Example Supplier", 250.0)], others=[b2], batch=got[b1["id"]])
check(st == "MATCH IN XERO" and "covers" in why, "a batched bill resolves with the payment noted")

# ---- VERIFY narrowed -----------------------------------------------------------

two = [line("2026-06-10", "Cafe One", 75.0), line("2026-06-10", "Cafe Two", 75.0)]
tb = lambda: bill("Example Catering Ltd", "2026-06-10", 75.0)   # noqa: E731
st, why, _ = run(tb(), two)
check(st == "VERIFY" and "2 free" in why, "two free lines of the amount, nothing between them: VERIFY")
st, why, _ = run(tb(), two, spends=[spend("2026-06-10", 75.0)])
check(st == "UNPAID" and "belong to other bills" in why, "every exact line already spoken for: UNPAID")
one_taken = [line("2026-06-10", "Cafe One", 75.0), line("2026-06-20", "Cafe Two", 75.0)]
st, _, _ = run(tb(), one_taken, spends=[spend("2026-06-10", 75.0)])
check(st == "MATCH IN XERO", "exactly one free line fits: resolved")
st, _, _ = run(tb(), [line("2026-06-10", "Cafe One", 75.0), line("2026-06-10", "Cafe One", 75.0)])
check(st == "MATCH IN XERO", "the same card line twice in an export is one payment")
st, why, _ = run(bill("Example Supplier Ltd", "2026-06-10", 480.0),
                 [line("2026-06-10", "Example Supplier", 450.0), line("2026-06-12", "Example Supplier", 500.0)])
check(st == "UNPAID" and "other charges" in why, "the supplier's other charges: UNPAID, not VERIFY")

# ---- the payables message --------------------------------------------------------

br.bank_history, br.bank_feed_lines = real_history, real_feed
today = date.today()
iso = lambda n: (today + timedelta(days=n)).isoformat()   # noqa: E731
u1 = bill("Example Supplier Ltd", iso(-40), 480.0, due=iso(-10), number="U-1", state="UNPAID", why="no payment")
u2 = bill("Acme Hosting", iso(-5), 90.0, ccy="EUR", due=iso(20), number="U-2", state="UNPAID", why="no payment")
md = bill("Example Rides", iso(-30), 40.0, due=iso(-20), number="M-1", state="MONTHLY DEBIT", why="monthly")
na = bill("Another Vendor", iso(-3), 55.0, status="DRAFT", number="N-1", state="NOT APPROVED",
          why="draft in Xero: approve it before paying, or void it")
vf = bill("Example Catering Ltd", iso(-9), 75.0, number="V-1", state="VERIFY", why="2 free bank lines fit. Pick one.")
store = {"entities": {KEY: {"bills": [u1, u2, md, na, vf], "spends": []}}, "assessed": "x"}

with tempfile.TemporaryDirectory() as tmp:
    br.RECONCILE_LIST = Path(tmp) / "reconcile_list.json"
    text = br.payables_message(store)
    check("*to pay*" in text and "total EUR 90.00 · USD 480.00" in text, "to pay carries totals per currency")
    check("Example Rides" not in text, "a MONTHLY DEBIT bill is never in the payer's list")
    check("*not approved in Xero*" in text and "Another Vendor" in text, "not approved section lists the draft")
    check("*check: may already be paid*" in text and "Example Catering" in text, "check section lists the VERIFY bill")
    check("the reconcile list has not been built" in text, "no reconcile list: one line says so")
    line_u1 = next(l for l in text.splitlines() if "U-1" in l)
    check(line_u1.endswith("U-1") and "overdue 10d" in line_u1, "bill line ends in the invoice number, shows overdue")
    br.RECONCILE_LIST.write_text(json.dumps({"built": datetime.now(timezone.utc).isoformat(), "items": [
        {"entity": KEY, "text": "reconcile the card line to the bill", "refs": "Acme Hosting · EUR 90.00 · U-2"}]}))
    text = br.payables_message(store)
    check("*to match in Xero*" in text and "reconcile the card line to the bill · Acme Hosting" in text,
          "reconcile list built today is read")
    br.RECONCILE_LIST.write_text(json.dumps({"built": "2020-01-01T00:00:00+00:00", "items": []}))
    check("not today" in br.payables_message(store), "a stale reconcile list is said to be stale")

    long_text = "\n".join(f"• line {i} " + "x" * 80 for i in range(200))
    parts = br._chunks(long_text)
    check(len(parts) > 1 and all(len(p) <= br.SLACK_CHUNK for p in parts), "long messages are chunked under the Slack limit")

    # the workbook: statuses in the legend, drafts on the entity tab, no overdue for a monthly debit
    for b in store["entities"][KEY]["bills"]:
        b["reconciled"] = False
    store["entities"][KEY].update({"pulled": "x", "from": "2025-01-01", "short_code": ""})
    store["counts"] = {KEY: {"UNPAID": 2, "MONTHLY DEBIT": 1, "NOT APPROVED": 1, "VERIFY": 1}}
    from openpyxl import load_workbook
    out = br.build_workbook(store, Path(tmp) / "bills.xlsx")
    wb = load_workbook(out)
    flat = [str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value]
    check("NOT APPROVED" in flat and "MONTHLY DEBIT" in flat, "workbook legend has the new statuses")
    tab = wb[config.entity(KEY).short]
    states = [tab.cell(r, 1).value for r in range(5, tab.max_row + 1)]
    check("NOT APPROVED" in states and "MONTHLY DEBIT" in states, "entity tab lists drafts and monthly debits")
    due_col = [tab.cell(r, 6).value for r in range(5, tab.max_row + 1)]
    check("monthly debit" in due_col, "a monthly debit bill is never overdue")

    top, detail = br.channel_report({"counts": store["counts"]}, "sent to the payer", ["Drive publish failed"])
    check(top.startswith(config.company_name().upper() + " ACCOUNTS PAYABLE") and "*failed*" in detail,
          "channel headline uses the company name; failures are listed")

# ---- the classifier tunables load from config ------------------------------------

check((br.NAMED_LEAD_DAYS, br.NAMED_LEAD_MIN, br.NAMED_AMOUNT_TOLERANCE) == (35, 1000.0, 0.01), "named defaults")
check((br.STRANGER_FX_TOLERANCE, br.STRANGER_DAYS) == (0.02, 4), "stranger defaults")

print()
if FAILS:
    print(f"{len(FAILS)} failed")
    sys.exit(1)
print("all passed")
