#!/usr/bin/env python
"""Bank and FX fees: every fee a group bank charged since the start date and
every currency conversion, pulled from the bank and kept up to date.

    data/bankfees/store.json                         the store, one row per movement
    data/reports/Bank and FX Fees Source Data.xlsx   the source data, rebuilt each refresh
    data/reports/Bank and FX Fees Analysis.xlsx      the analysis, rebuilt each refresh
    published to Drive (kind bank_fees):  <publish folder> / <bank_fees subfolder> /
                         both workbooks, overwritten every refresh

WHICH ACCOUNTS. Every bank account in config/group.toml whose provider has an
API (revolut or mercury), with its saved statement CSV when it has one. A
`csv` account has no API and no known export schema that carries fee detail,
so it is not part of this report.

THE ANALYSIS. One tab, Analysis: every fee in the reporting currency
([company] reporting_currency) by month, one stacked bar per
bank provider (every entity's Revolut accounts together, every entity's
Mercury accounts together), stacked by fee type, with the table the chart is
drawn from. A fee in another currency is converted at its month's average ECB
reference rate, held in the store and shown on the Rates tab: approximate on
purpose.

THE TARIFF. What each fee is made of is read off its size against the price
list in config/group.toml [bank_fees.<provider>], filled from your own bank
plan. The part names (FX markup, weekend FX markup, guaranteed SWIFT, ACH,
international transfer, plan charge) are generic examples of the fees a
business bank commonly charges, not a statement of what any plan includes: a
part your plan does not charge, or that has no price configured, is simply
never matched. A provider with no tariff configured has its fees collected and
charted unsplit, as one "fee, no tariff configured" type, and no expected
price is checked.

WHAT IS KEPT. Only movements that carry a fee or convert a currency. Each row
has a kind, what the movement is:

    conversion            a Revolut EXCHANGE between two of the entity's own
                          accounts, both sides and the fee on it
    foreign payment       a payment or receipt in a currency the account does
                          not hold: the bank converts inside the payment (a
                          Revolut transfer or card payment, a Mercury FX wire,
                          a Mercury card charge in another currency)
    payment, card payment a Revolut transfer or card payment in the account's
                          own currency that carries a fee
    plan fee              the Revolut Business plan charge (type FEE / charge)
    card intl fee         Mercury's Intl. Transaction Fee, its own line
    wire fx fee           Mercury's Intl. Wire Foreign Exchange Fee, its own line
    wire fee              Mercury's Intl. Wire Fee, its own line

and a row with a fee on it a fee type as well: conversion fee, payment fee or
card fee for a Revolut fee charged on the movement, else the kind itself.

Fees are positive when charged, negative when refunded. A Revolut fee sits on
the transaction it was charged on, so a conversion or foreign payment row can
also be a fee row; a Mercury fee is always its own line, and its FX wire
carries the linked fee only as information, so the Fees sheet never counts one
fee twice.

WHERE IT COMES FROM (docs/BANKING.md). Per account, the saved statement CSV for
movements completed up to and including its `csv_until`, the read-only bank
API from the day after (`api_from`); an account with no `csv_until` is read
from the API only. Set `csv_until` on the last WHOLE day an export holds (an
export taken part-way through a day leaves that day to the API). A movement is
dated by the day it settled on both sides of the cut-over, so the two never
overlap and nothing falls between them.

HOW THE REFRESH STAYS COMPLETE WITHOUT DUPLICATES. Each account carries a
`through` date, the day it was last pulled to. A refresh pulls the API from
OVERLAP_DAYS before that date to today and upserts by the bank's own
transaction id, so a movement read twice is the same row, a payment that was
pending last time and has settled since is picked up, and one that has since
been reverted or declined is dropped. Never a pull of the whole period: the
exports are read once, when the store is first seeded.

THE START DATE is `--from YYYY-MM-DD`, else [company].records_from in
config/group.toml, else the start of the previous financial year. A store
seeded from a different start date is rebuilt.

SCHEDULED: the `balance-reports` job in deploy/run-scheduled.sh runs
`bank_fees.py refresh` every day it runs; run it by hand on the server to
refresh the workbooks in between.

READ-ONLY against the banks (both clients refuse every non-GET) and never
touches Xero. The only other read is the ECB rates from frankfurter.dev. Runs
on the server, where the bank credentials are.

Usage:
    .venv/bin/python scripts/bank_fees.py refresh [--from YYYY-MM-DD] [--no-publish]   # top up, rebuild, publish
    .venv/bin/python scripts/bank_fees.py build [--no-publish]     # rebuild both workbooks from the store
    .venv/bin/python scripts/bank_fees.py status
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from accounting_agent import config  # noqa: E402

STORE = ROOT / "data" / "bankfees" / "store.json"
STATEMENTS = ROOT / "data" / "statements"
SOURCE_BOOK = ROOT / "data" / "reports" / "Bank and FX Fees Source Data.xlsx"
ANALYSIS_BOOK = ROOT / "data" / "reports" / "Bank and FX Fees Analysis.xlsx"

OVERLAP_DAYS = 21                  # re-read window: pending settles, late reversals
BANK_NAMES = {"revolut": "Revolut", "mercury": "Mercury"}
BANKS = ("Revolut", "Mercury")


def _accounts() -> dict[str, dict]:
    """Store key -> {short, bank, slug, csv, csv_until, api_from, currency}, for
    every configured revolut / mercury account, in config order."""
    out = {}
    for e in config.entities():
        for b in e.bank_accounts:
            if not b.has_api:
                continue
            out[f"{e.slug}:{b.provider}:{b.slug}"] = {
                "short": e.short, "bank": BANK_NAMES[b.provider], "slug": b.slug,
                "csv": b.statement_csv if b.csv_until else "", "csv_until": b.csv_until,
                "api_from": b.api_from, "currency": b.currency,
            }
    return out


def _start(store: dict) -> str:
    return store.get("from") or config.records_from()


# Mercury's currencyExchangeInfo names the sold side "converted" + "From...".
MERCURY_FX_FROM = "converted" "From"

MERCURY_FEES = {
    "Intl. Transaction Fee": "card intl fee",
    "Intl. Wire Foreign Exchange Fee": "wire fx fee",
    "Intl. Wire Fee": "wire fee",
}

# what a Source Data row is, for the Coverage sheet
SOURCE_KINDS = [
    ("conversion", "Revolut exchange between two of the entity's own accounts; fee on the exchange"),
    ("foreign payment", "payment or receipt in a currency the account does not hold, converted inside it"),
    ("payment", "Revolut transfer in the account's own currency that carries a fee"),
    ("card payment", "Revolut card payment in the account's own currency that carries a fee"),
    ("plan fee", "Revolut Business plan charge"),
    ("card intl fee", "Mercury international card fee, its own line"),
    ("wire fx fee", "Mercury foreign exchange fee on an international wire, its own line"),
    ("wire fee", "Mercury international wire fee, its own line"),
]

# kinds that are a fee and nothing else: the row is the fee line itself
FEE_LINES = {"plan fee", "card intl fee", "wire fx fee", "wire fee"}

COLUMNS = [
    ("date", "Date"), ("month", "Month"), ("entity", "Entity"), ("bank", "Bank"),
    ("account", "Account"), ("kind", "Kind"), ("bank_type", "Bank type"),
    ("description", "Description"), ("cardholder", "Card / payer"),
    ("amount_ccy", "Amount CCY"), ("amount", "Amount"),
    ("sold_ccy", "Sold CCY"), ("sold", "Sold"), ("bought_ccy", "Bought CCY"),
    ("bought", "Bought"), ("rate", "Rate (bought per sold)"),
    ("fee_type", "Fee type"), ("fee_ccy", "Fee CCY"), ("fee", "Fee"), ("linked_fee", "Linked fee line"),
    ("txn_id", "Transaction id"), ("source", "Source"),
]


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _row(**kw) -> dict:
    row = {k: None for k, _ in COLUMNS}
    row.update(kw)
    row["month"] = row["date"][:7]
    if row.get("fee"):
        typ = (row.get("bank_type") or "").upper()
        row["fee_type"] = (row["kind"] if row["kind"] in FEE_LINES
                           else "conversion fee" if typ == "EXCHANGE"
                           else "card fee" if typ.startswith("CARD") else "payment fee")
    if row.get("sold") and row.get("bought"):
        row["rate"] = round(row["bought"] / row["sold"], 6)
    if row.get("amount") is None and row.get("sold"):
        # a conversion's size is what it sold
        row["amount"], row["amount_ccy"] = row["sold"], row["sold_ccy"]
    return row


# ---------------------------------------------------------------- exports

def _revolut_csv(acc: dict, start: str) -> dict[str, dict]:
    short, bank, end = acc["short"], acc["bank"], acc["csv_until"]
    path = STATEMENTS / acc["csv"]
    if not acc["csv"] or not path.exists():
        return {}
    rows = [r for r in csv.DictReader(open(path, newline="", encoding="utf-8-sig"))
            if r["State"] == "COMPLETED"
            and start <= (r["Date completed (UTC)"] or "")[:10] <= end
            and (not acc["currency"] or r["Payment currency"] == acc["currency"])]
    out: dict[str, dict] = {}
    exchanges: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r["Type"] == "EXCHANGE":
            exchanges[r["ID"]].append(r)
            continue
        date, ccy, amount = r["Date completed (UTC)"][:10], r["Payment currency"], _num(r["Amount"])
        common = dict(date=date, entity=short, bank=bank, account=r["Account"],
                      amount=abs(amount), amount_ccy=ccy,
                      bank_type=r["Type"], description=r["Description"],
                      cardholder=r["Card label"] or r["Payer"] or None,
                      txn_id=r["ID"], source="statement export")
        if r["Type"] == "FEE":
            out[f"rv:{r['ID']}:{ccy}:{'-' if amount < 0 else '+'}"] = _row(
                kind="plan fee", fee_ccy=ccy, fee=round(-amount, 2), **common)
            continue
        fee = round(-_num(r["Fee"]), 2)
        orig = r["Orig currency"]
        foreign = bool(orig) and orig != ccy
        if not fee and not foreign:
            continue
        row = _row(kind=("foreign payment" if foreign else
                         "card payment" if r["Type"].startswith("CARD") else "payment"),
                   fee_ccy=(r["Fee currency"] or ccy) if fee else None, fee=fee or None,
                   **common)
        if foreign:
            mine, theirs = abs(amount), abs(_num(r["Orig amount"]))
            if amount < 0:
                row.update(sold_ccy=ccy, sold=mine, bought_ccy=orig, bought=theirs)
            else:
                row.update(sold_ccy=orig, sold=theirs, bought_ccy=ccy, bought=mine)
            row = _row(**row)
        out[f"rv:{r['ID']}:{ccy}:{'-' if amount < 0 else '+'}"] = row
    for tid, legs in exchanges.items():
        sold = next((l for l in legs if _num(l["Amount"]) < 0), None)
        bought = next((l for l in legs if _num(l["Amount"]) > 0), None)
        if not sold or not bought:
            continue
        fee = round(-_num(sold["Fee"]), 2)
        out[f"rv:{tid}"] = _row(
            date=sold["Date completed (UTC)"][:10], entity=short, bank=bank,
            account=f"{sold['Account']} → {bought['Account']}", kind="conversion",
            bank_type="EXCHANGE", description=sold["Description"],
            sold_ccy=sold["Payment currency"], sold=abs(_num(sold["Amount"])),
            bought_ccy=bought["Payment currency"], bought=abs(_num(bought["Amount"])),
            fee_ccy=(sold["Fee currency"] or None) if fee else None, fee=fee or None,
            txn_id=tid, source="statement export")
    return out


def _mercury_account(name: str) -> str:
    return (name or "").replace("••", "xx")


def _mercury_csv_key(r: dict, seen: Counter) -> str:
    # the export carries no transaction id: the timestamp, the amount, the
    # description and the account name the line, and a counter separates two
    # identical lines in the same second
    base = "|".join((r["Timestamp"], r["Amount"], r["Description"], r["Source Account"]))
    seen[base] += 1
    return "mc-csv:" + hashlib.sha1(f"{base}|{seen[base]}".encode()).hexdigest()[:16]


# Mercury holds US dollar accounts only, so its own movements and fees are USD
# whatever the group's reporting currency.
MERCURY_CCY = "USD"


def _mercury_csv(acc: dict, start: str) -> dict[str, dict]:
    short, bank, end = acc["short"], acc["bank"], acc["csv_until"]
    path = STATEMENTS / acc["csv"]
    if not acc["csv"] or not path.exists():
        return {}
    out: dict[str, dict] = {}
    seen: Counter = Counter()
    for r in csv.DictReader(open(path, newline="", encoding="utf-8-sig")):
        m, d, y = r["Date (UTC)"].split("-")
        date = f"{y}-{m}-{d}"
        # the export dates a sent line by the day it posted, as the API does
        if r["Status"] != "Sent" or not (start <= date <= end):
            continue
        k = _mercury_csv_key(r, seen)
        amount = _num(r["Amount"])
        common = dict(date=date, entity=short, bank=bank,
                      account=_mercury_account(r["Source Account"]),
                      amount=abs(amount), amount_ccy=MERCURY_CCY,
                      bank_type=r["Merchant Type"] or r["Category"] or None,
                      description=r["Description"], cardholder=r["Name On Card"] or None,
                      txn_id=None, source="statement export")
        if r["Description"] in MERCURY_FEES:
            out[k] = _row(kind=MERCURY_FEES[r["Description"]], fee_ccy=MERCURY_CCY,
                          fee=round(-amount, 2), **common)
        elif r["Original Currency"] and r["Original Currency"] != MERCURY_CCY:
            # the export names the currency the merchant charged in but not
            # the amount in it
            out[k] = _row(kind="foreign payment", sold_ccy=MERCURY_CCY, sold=abs(amount),
                          bought_ccy=r["Original Currency"], **common)
    return out


# ---------------------------------------------------------------- APIs

def _api_start(acc: dict, start: str) -> str:
    return max(acc["api_from"] or start, start)


def _revolut_api(acc: dict, start: str, end: str, floor: str) -> tuple[dict[str, dict], set[str]]:
    """(rows to keep, keys the bank now says are not cash). `floor` is the
    first date the API is the source for this account."""
    from accounting_agent.revolut import RevolutClient
    short, bank = acc["short"], acc["bank"]
    client = RevolutClient(acc["slug"])
    names = {}
    for a in client.accounts():
        names[a["id"]] = config.statement_account_name("revolut", a.get("currency"), a.get("name"))
    keep: dict[str, dict] = {}
    gone: set[str] = set()
    for t in client.transactions(start, end):
        legs = t.get("legs") or []
        if acc["currency"]:
            legs = [l for l in legs if l.get("currency") == acc["currency"]] if t.get("type", "").upper() != "EXCHANGE" else legs
        date = (t.get("completed_at") or t.get("created_at") or "")[:10]
        completed = t.get("state") == "completed" and date >= floor
        typ = (t.get("type") or "").upper()
        if typ == "EXCHANGE":
            sold = next((l for l in legs if _num(l.get("amount")) < 0), None)
            bought = next((l for l in legs if _num(l.get("amount")) > 0), None)
            k = f"rv:{t['id']}"
            if not completed or not sold or not bought:
                gone.add(k)
                continue
            fee_leg = next((l for l in legs if l.get("fee")), None)
            fee = round(_num(fee_leg.get("fee")), 2) if fee_leg else 0
            keep[k] = _row(
                date=date, entity=short, bank=bank,
                account=f"{names.get(sold['account_id'], sold.get('currency'))} → "
                        f"{names.get(bought['account_id'], bought.get('currency'))}",
                kind="conversion", bank_type=typ, description=sold.get("description"),
                sold_ccy=sold.get("currency"), sold=abs(_num(sold.get("amount"))),
                bought_ccy=bought.get("currency"), bought=abs(_num(bought.get("amount"))),
                fee_ccy=fee_leg.get("currency") if fee else None, fee=fee or None,
                txn_id=t["id"], source="Revolut API")
            continue
        card = t.get("card") or {}
        for leg in legs:
            ccy, amount = leg.get("currency"), _num(leg.get("amount"))
            k = f"rv:{t['id']}:{ccy}:{'-' if amount < 0 else '+'}"
            if not completed:
                gone.add(k)
                continue
            common = dict(date=date, entity=short, bank=bank,
                          account=names.get(leg.get("account_id"), ccy), bank_type=typ,
                          amount=abs(amount), amount_ccy=ccy,
                          description=leg.get("description") or t.get("reference"),
                          cardholder=" ".join(x for x in (card.get("first_name"), card.get("last_name")) if x) or None,
                          txn_id=t["id"], source="Revolut API")
            if typ in ("FEE", "CHARGE"):
                keep[k] = _row(kind="plan fee", fee_ccy=ccy, fee=round(-amount, 2), **common)
                continue
            fee = round(_num(leg.get("fee")), 2)
            bill_ccy = leg.get("bill_currency")
            foreign = bool(bill_ccy) and bill_ccy != ccy
            if not fee and not foreign:
                continue
            row = dict(kind=("foreign payment" if foreign else
                             "card payment" if typ.startswith("CARD") else "payment"),
                       fee_ccy=ccy if fee else None, fee=fee or None, **common)
            if foreign:
                mine, theirs = abs(amount), abs(_num(leg.get("bill_amount")))
                if amount < 0:
                    row.update(sold_ccy=ccy, sold=mine, bought_ccy=bill_ccy, bought=theirs)
                else:
                    row.update(sold_ccy=bill_ccy, sold=theirs, bought_ccy=ccy, bought=mine)
            keep[k] = _row(**row)
    return keep, gone


def _mercury_api(acc: dict, start: str, end: str, floor: str) -> tuple[dict[str, dict], set[str]]:
    from accounting_agent.mercury import MercuryClient
    short, bank = acc["short"], acc["bank"]
    client = MercuryClient(acc["slug"])
    keep: dict[str, dict] = {}
    gone: set[str] = set()
    for a in client.accounts():
        for t in client.transactions(a["id"], start, end):
            k = f"mc:{t['id']}"
            date = (t.get("postedAt") or "")[:10]
            if t.get("status") != "sent" or not date or date < floor:
                gone.add(k)
                continue
            kind, amount = t.get("kind") or "", _num(t.get("amount"))
            desc = t.get("bankDescription") or t.get("counterpartyName")
            common = dict(date=date, entity=short, bank=bank,
                          account=_mercury_account(a.get("nickname") or a.get("name")),
                          amount=abs(amount), amount_ccy=MERCURY_CCY,
                          bank_type=kind, description=desc, txn_id=t["id"], source="Mercury API")
            fx = t.get("currencyExchangeInfo")
            if kind == "cardInternationalTransactionFee":
                keep[k] = _row(kind="card intl fee", fee_ccy=MERCURY_CCY, fee=round(-amount, 2), **common)
            elif kind == "wireFee":
                fx_fee = "exchange" in (desc or "").lower()
                keep[k] = _row(kind="wire fx fee" if fx_fee else "wire fee", fee_ccy=MERCURY_CCY,
                               fee=round(-amount, 2), **common)
            elif fx:
                keep[k] = _row(kind="foreign payment",
                               sold_ccy=fx.get(MERCURY_FX_FROM + "Currency"),
                               sold=abs(_num(fx.get(MERCURY_FX_FROM + "Amount"))),
                               bought_ccy=fx.get("convertedToCurrency"),
                               bought=abs(_num(fx.get("convertedToAmount"))),
                               linked_fee=_num(fx.get("feeAmount")) or None, **common)
    return keep, gone


# ---------------------------------------------------------------- store

# The shape of a stored row. A store of an older shape, or seeded from another
# start date, is rebuilt from scratch on the next refresh: the exports again
# and the API from the cut-over, the one time the whole period is read after
# the first seed.
SCHEMA = 3      # 3: accounts keyed entity:provider:slug, start date on the store


def load_store(start: str | None = None) -> dict:
    if STORE.exists():
        store = json.loads(STORE.read_text())
        if store.get("schema") == SCHEMA and (not start or store.get("from") == start):
            return store
    return {"schema": SCHEMA, "from": start or "", "rows": {}, "through": {}, "seeded": {}, "last_refresh": {}}


def save_store(store: dict) -> None:
    STORE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STORE.with_suffix(".tmp")
    tmp.write_text(json.dumps(store, indent=1, sort_keys=True, default=str))
    tmp.replace(STORE)


def refresh(store: dict) -> dict:
    """Seed from the exports once, then top every account up from the API.
    Returns what changed, per account."""
    start = _start(store)
    store["from"] = start
    today = dt.date.today().isoformat()
    rows = store["rows"]
    changes: dict[str, dict] = {}
    for key, acc in _accounts().items():
        added = updated = removed = 0
        if not store["seeded"].get(key):
            seed = {}
            if acc["csv"]:
                seed = _revolut_csv(acc, start) if acc["bank"] == "Revolut" else _mercury_csv(acc, start)
            for k, row in seed.items():
                rows[k] = row
            added += len(seed)
            store["seeded"][key] = (f"exports {start} to {acc['csv_until']}, {len(seed)} rows" if acc["csv"]
                                    else "no export, API only")
        # both APIs filter on the date a movement was created, and one created
        # before the window can settle inside it, so the window always reaches
        # OVERLAP_DAYS back; rows dated before the API start are the exports' own
        floor = _api_start(acc, start)
        base = dt.date.fromisoformat(store["through"].get(key) or floor)
        pull_from = (base - dt.timedelta(days=OVERLAP_DAYS)).isoformat()
        end = (dt.date.today() + dt.timedelta(days=1)).isoformat()
        keep, gone = (_revolut_api if acc["bank"] == "Revolut" else _mercury_api)(acc, pull_from, end, floor)
        for k, row in keep.items():
            if k not in rows:
                added += 1
            elif rows[k] != row:
                updated += 1
            rows[k] = row
        for k in gone - set(keep):
            if k in rows and rows[k]["source"] != "statement export":
                del rows[k]
                removed += 1
        store["through"][key] = today
        changes[f"{acc['short']} {acc['bank']}"] = {"added": added, "updated": updated, "removed": removed,
                                                    "pulled": f"{pull_from} to {today}"}
    return changes


# ---------------------------------------------------------------- workbooks

def _sorted_rows(store: dict, pick) -> list[dict]:
    return sorted((r for r in store["rows"].values() if pick(r)),
                  key=lambda r: (r["date"], r["entity"], r["kind"], r.get("txn_id") or ""))


# ---------------------------------------------------------------- fee breakdown

# What each fee is made of. The statement carries one fee figure per movement;
# the parts are read off its size against the amount, using the price list the
# group configured for that provider in config/group.toml. The part names are
# generic examples of common business-bank fees, and the figures placeholders,
# not any bank's actual tariff:
#
#   [bank_fees.revolut]
#   tariff_currency = "GBP"   # currency the fixed prices below are quoted in
#   fx_pct = 0.6              # % on a conversion over the fee-free allowance
#   weekend_fx_pct = 1.0      # % extra on a conversion at the weekend
#   swift_fixed = 20.0        # a guaranteed (OUR, "pay all fees") SWIFT transfer
#   intl_fixed = 5.0          # a standard international transfer outside the allowance
#   ach_pct = 0.2             # % on a domestic ACH payment
#   ach_min = 0.50            # the ACH minimum, in tariff_currency
#
# Every key is optional; a part with no price is never checked, and a provider
# with no block at all has its fees charted as "fee, no tariff configured". A
# fixed price is converted into the fee's currency at the month's rate. A
# payment in another currency can carry the FX parts and a transfer charge
# together. A part is recognised within FEE_TOLERANCE of its price, because a
# fixed fee charged in another currency is converted at the day's rate and the
# check uses the month's; what matches nothing is "unexplained".
FEE_PARTS = [
    ("fx", "FX fee, over the allowance"),
    ("fx_part", "FX fee on part of the amount"),
    ("weekend", "weekend FX fee"),
    ("swift", "guaranteed SWIFT fee"),
    ("ach", "ACH fee"),
    ("intl", "international transfer fee"),
    ("plan", "Revolut plan fee"),
    ("card3", "Mercury card FX fee"),
    ("wire_fx", "Mercury wire FX fee"),
    ("wire", "Mercury wire fee"),
    ("untariffed", "fee, no tariff configured"),
    ("refund", "fee refunded"),
    ("other", "unexplained"),
]
FEE_TOLERANCE = 0.04


def tariff(bank: str) -> dict:
    """The configured price list for one provider ([bank_fees.<provider>]);
    {} when the group has not filled one in."""
    try:
        return config.bank_fee_tariffs().get(bank.lower()) or {}
    except config.ConfigError:
        return {}


def fee_parts(store: dict, row: dict) -> dict[str, float]:
    """{part: amount in the fee's currency}; the parts add up to the fee."""
    fee, ccy, kind = row.get("fee") or 0, row.get("fee_ccy"), row["kind"]
    if not fee:
        return {}
    if fee < 0:
        return {"refund": fee}
    if kind in ("plan fee", "card intl fee", "wire fx fee", "wire fee"):
        return {{"plan fee": "plan", "card intl fee": "card3", "wire fx fee": "wire_fx",
                 "wire fee": "wire"}[kind]: fee}
    t = tariff(row.get("bank") or "")
    if not t:
        return {"untariffed": fee}
    tccy = str(t.get("tariff_currency") or ccy or "").upper()

    def pct(key: str) -> float | None:
        v = t.get(key)
        return float(v) / 100 if v not in (None, "", 0) else None

    def fixed(key: str) -> float | None:
        """A fixed price from the tariff, in the fee's currency."""
        v = t.get(key)
        if v in (None, "", 0):
            return None
        if tccy == ccy:
            return float(v)
        in_rep = _to_rep(store, row["month"], tccy, float(v))
        per_unit = _to_rep(store, row["month"], ccy, 1.0)
        return in_rep / per_unit if in_rep and per_unit else None

    swift, intl, ach_min = fixed("swift_fixed"), fixed("intl_fixed"), fixed("ach_min")
    fx_pct, wk_pct, ach_pct = pct("fx_pct"), pct("weekend_fx_pct"), pct("ach_pct")

    def near(value: float, target: float | None, tol: float = FEE_TOLERANCE) -> bool:
        return bool(target) and abs(value - target) <= max(tol * target, 0.015)

    if kind in ("conversion", "foreign payment"):
        base = (row.get("sold") if ccy == row.get("sold_ccy")
                else row.get("bought") if ccy == row.get("bought_ccy") else None)
        if not base:
            return {"other": fee}
        fx = fx_pct * base if fx_pct else None
        wk = wk_pct * base if wk_pct else None
        if near(fee, fx, 0.01):
            return {"fx": fee}
        if near(fee, wk, 0.01):
            return {"weekend": fee}
        if fx and wk and near(fee, fx + wk, 0.01):
            return {"fx": round(fee * fx / (fx + wk), 2), "weekend": round(fee * wk / (fx + wk), 2)}
        # a transfer fee with or without the FX parts (none while the movement
        # was inside the allowance): the closest combination wins
        extras = [{}]
        if fx:
            extras.append({"fx": fx})
        if wk:
            extras.append({"weekend": wk})
        if fx and wk:
            extras.append({"fx": fx, "weekend": wk})
        best, err = None, None
        for fixed_part, price in (("swift", swift), ("intl", intl)):
            if not price:
                continue
            for extra in extras:
                total = price + sum(extra.values())
                # tolerance on the fixed charge alone: on the whole fee a large
                # FX part would swallow it
                if abs(fee - total) <= max(FEE_TOLERANCE * price, 0.015) \
                        and (err is None or abs(fee - total) < err):
                    err = abs(fee - total)
                    best = {k: round(v, 2) for k, v in extra.items()}
                    best[fixed_part] = round(fee - sum(best.values()), 2)
        if best:
            return best
        if fx and ach_pct and near(fee, fx + ach_pct * base, 0.01):
            # converted, then sent to a domestic bank: the conversion and the ACH charge
            return {"fx": round(fx, 2), "ach": round(fee - round(fx, 2), 2)}
        if fx and fee < fx:
            return {"fx_part": fee}
        if fx:
            return {"fx": round(fx, 2), "other": round(fee - fx, 2)}
        return {"other": fee}
    amount = row.get("amount") or 0
    for part, price in (("swift", swift), ("ach", ach_pct * amount if ach_pct else None), ("intl", intl)):
        if price and near(fee, price, 0.01 if part == "ach" else FEE_TOLERANCE):
            return {part: fee}
    if ach_min and ach_pct and amount and ach_pct * amount < fee and near(fee, ach_min):
        return {"ach": fee}      # the ACH minimum on a small payment
    return {"other": fee}


def _cutover_text() -> str:
    parts = []
    for acc in _accounts().values():
        if acc["csv"]:
            parts.append(f"{acc['short']} {acc['bank']}: export to {acc['csv_until']}, API after")
        else:
            parts.append(f"{acc['short']} {acc['bank']}: API")
    return "; ".join(parts) or "no revolut or mercury account configured"


def build_source(store: dict, out: Path = SOURCE_BOOK) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="203864")
    start = _start(store)
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"
    ws["A1"] = f"Bank and FX fees, source data, from {start}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Every fee the group's banks charged and every currency conversion, {start} to "
                f"{min(store['through'].values(), default='')}. {_cutover_text()}. "
                f"Refreshed {store['last_refresh'].get('at', '')} UTC. Nothing here comes from Xero.")
    ws["A2"].font = Font(size=9, italic=True, color="555555")
    hdr = ["Entity", "Bank", "Exports", "API pulled through", "Fee rows", "Conversion rows",
           "Fees by currency"]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(4, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    r = 5
    for key, acc in _accounts().items():
        mine = [x for x in store["rows"].values() if x["entity"] == acc["short"] and x["bank"] == acc["bank"]]
        fees: dict[str, float] = defaultdict(float)
        for x in mine:
            if x.get("fee"):
                fees[x["fee_ccy"]] += x["fee"]
        vals = [acc["short"], acc["bank"], store["seeded"].get(key, ""), store["through"].get(key, ""),
                sum(1 for x in mine if x.get("fee")),
                sum(1 for x in mine if x["kind"] in ("conversion", "foreign payment")),
                ", ".join(f"{c} {v:,.2f}" for c, v in sorted(fees.items()))]
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, v)
        r += 1
    r += 1
    ws.cell(r, 1, "Kinds").font = Font(bold=True)
    for kind, what in SOURCE_KINDS:
        r += 1
        ws.cell(r, 1, kind).font = Font(bold=True)
        ws.cell(r, 2, what)
    r += 2
    for note in (
        "Fees are positive when charged, negative when refunded, in the currency the bank charged them in.",
        "Fee type says what the fee was for: conversion fee, payment fee or card fee for a Revolut fee on the movement, else the kind.",
        "A Mercury FX wire shows its fee under Linked fee line for reference; the fee itself is its own wire fx fee row, so Fees counts it once.",
        "A Mercury card charge in another currency: the export names the currency, not the amount in it; the API does not mark it at all, its international card fee line does.",
        "Rate is bought per sold as the bank booked it, after any fee taken from the bought side.",
    ):
        ws.cell(r, 1, note).font = Font(size=9, color="555555")
        r += 1
    ws.column_dimensions["A"].width = 18
    for col, w in zip("BCDEFG", (10, 40, 18, 10, 14, 50)):
        ws.column_dimensions[col].width = w

    for title, pick in (
        ("Fees", lambda x: bool(x.get("fee"))),
        ("Conversions", lambda x: x["kind"] in ("conversion", "foreign payment")),
    ):
        sh = wb.create_sheet(title)
        # the Fees sheet carries the fee's parts after the row's own columns
        parts = FEE_PARTS if title == "Fees" else []
        labels = [label for _, label in COLUMNS] + ["Fee breakdown"] * bool(parts) + [p for _, p in parts]
        for j, label in enumerate(labels, start=1):
            c = sh.cell(1, j, label)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(_sorted_rows(store, pick), start=2):
            for j, (k, _) in enumerate(COLUMNS, start=1):
                c = sh.cell(i, j, row.get(k))
                if k in ("sold", "bought", "fee", "linked_fee", "amount"):
                    c.number_format = "#,##0.00"
            if parts:
                got = fee_parts(store, row)
                names = dict(FEE_PARTS)
                sh.cell(i, len(COLUMNS) + 1,
                        " + ".join(f"{names[p]} {v:,.2f}" for p, v in got.items()))
                for j, (p, _) in enumerate(parts, start=len(COLUMNS) + 2):
                    if got.get(p):
                        sh.cell(i, j, got[p]).number_format = "#,##0.00"
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{sh.max_row}"
        for j, (k, _) in enumerate(COLUMNS, start=1):
            sh.column_dimensions[get_column_letter(j)].width = {
                "description": 36, "account": 22, "txn_id": 38, "cardholder": 20}.get(k, 12)
        if parts:
            sh.column_dimensions[get_column_letter(len(COLUMNS) + 1)].width = 44
            for j in range(len(COLUMNS) + 2, len(labels) + 1):
                sh.column_dimensions[get_column_letter(j)].width = 13
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# The analysis: a month by month stacked bar, one bar per bank provider, each
# stacked by fee type, in the reporting currency, then a section per fee type.
# The fee types are the parts fee_parts() reads off each fee, grouped: FX
# whether the whole movement or only part of it was over the allowance; the
# international transfer fee, the Mercury wire fee, unexplained and refunded
# fees are left out. A fixed order and colour (dataviz categorical slots 1-8),
# so a type keeps its colour whatever a month holds.
KINDS = [
    ("FX fee", ("Revolut",), "2A78D6",
     "the FX fee on conversions and payments in another currency above the plan's fee-free "
     "allowance, including a movement only partly over it", ("fx", "fx_part")),
    ("Revolut plan fee", ("Revolut",), "EB6834", "the monthly Revolut Business plan charge", ("plan",)),
    ("ACH fee", ("Revolut",), "1BAF7A",
     "the transfer fee on a domestic ACH payment, where the configured tariff prices one",
     ("ach",)),
    ("Weekend FX fee", ("Revolut",), "EDA100", "the extra FX fee on a conversion at the weekend",
     ("weekend",)),
    ("Mercury card FX", ("Mercury",), "E87BA4", "Mercury's fee on card spend in another currency", ("card3",)),
    ("Mercury wire FX", ("Mercury",), "008300",
     "Mercury's FX fee on an international wire in another currency", ("wire_fx",)),
    ("SWIFT fee", ("Revolut",), "4A3AA7",
     "the guaranteed SWIFT (\"pay all fees\") option, so the recipient gets the full amount; "
     "charged in the account's currency", ("swift",)),
    ("Fee, no tariff", ("Revolut",), "7F7F7F",
     "every fee from a provider with no [bank_fees] price list in config/group.toml, not split "
     "into its parts", ("untariffed",)),
]
PART_KIND = {part: kind for kind, _, _, _, parts in KINDS for part in parts}
ANALYSIS_WIDTHS = {"A": 13, "B": 8, "C": 9, "D": 13.57, "E": 14.57}


def reporting_ccy() -> str:
    return str(config.reporting_currency()).upper()


def monthly_rates(currencies: set[str], start: str, end: str) -> dict[str, dict[str, float]]:
    """{month: {ccy: reporting currency per unit}}, the average of the month's
    ECB reference rates (frankfurter, the source scripts/bills_report.py uses).
    Approximate by design: a fee is converted at its month's average, not at
    the rate of its own day."""
    import urllib.parse
    import urllib.request
    base = reporting_ccy()
    wanted = sorted(c for c in currencies if c and c != base)
    if not wanted:
        return {}
    url = (f"https://api.frankfurter.dev/v1/{start}..{end}?"
           + urllib.parse.urlencode({"base": base, "symbols": ",".join(wanted)}))
    req = urllib.request.Request(url, headers={"Accept": "application/json",
                                               "User-Agent": "accounting-agent-bank-fees/1.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        daily = json.load(resp).get("rates") or {}
    sums: dict[str, dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for day, quotes in daily.items():
        for ccy, per_base in quotes.items():
            if per_base:
                sums[day[:7]][ccy].append(1 / per_base)
    # a range opening on a holiday starts at the fixing before it
    return {m: {c: sum(v) / len(v) for c, v in q.items()} for m, q in sums.items()
            if m >= start[:7]}


def _to_rep(store: dict, month: str, ccy: str | None, value: float | None) -> float | None:
    """`value` in `ccy` converted to the reporting currency at the month's rate."""
    if value is None or not ccy:
        return None
    if ccy == reporting_ccy():
        return value
    if store.get("fx_base") != reporting_ccy():
        return None                 # rates held are for another currency; refresh_rates replaces them
    months = store.get("fx") or {}
    rate = (months.get(month) or {}).get(ccy)
    if rate is None:
        # the running month before its first fixing: the latest month there is
        earlier = [m for m in sorted(months) if m <= month and ccy in months[m]]
        rate = months[earlier[-1]][ccy] if earlier else None
    return value * rate if rate is not None else None


def _who_banks_where() -> str:
    """'Revolut is A and B together; Mercury is C.' from config."""
    by_bank: dict[str, list[str]] = defaultdict(list)
    for acc in _accounts().values():
        if acc["short"] not in by_bank[acc["bank"]]:
            by_bank[acc["bank"]].append(acc["short"])
    parts = [f"{b} is {', '.join(s)}{' together' if len(s) > 1 else ''}" for b, s in by_bank.items()]
    return "; ".join(parts) + "." if parts else ""


def build_analysis(store: dict, out: Path = ANALYSIS_BOOK) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.data_source import AxDataSource, StrRef
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Alignment, Font, PatternFill

    head = PatternFill("solid", fgColor="203864")
    grey = Font(size=9, italic=True, color="555555")
    rows = list(store["rows"].values())
    months = sorted({r["month"] for r in rows})
    # the bars: only the providers the group actually banks with
    present = {acc["bank"] for acc in _accounts().values()} | {r["bank"] for r in rows}
    banks = [b for b in BANKS if b in present] or list(BANKS)
    rep = reporting_ccy()
    used_parts = {p for r in rows for p in fee_parts(store, r)}
    kinds = [k for k in KINDS if any(b in banks for b in k[1])
             and (set(k[4]) & used_parts or (k[4] != ("untariffed",) and any(tariff(b) for b in k[1])))]
    series_def = [(bank, kind, colour) for kind, kb, colour, _, _ in kinds for bank in kb if bank in banks]
    fee_rep: dict[tuple, float] = defaultdict(float)     # (month, bank, fee type)
    count: Counter = Counter()
    unpriced: Counter = Counter()
    for r in rows:
        for part, value in fee_parts(store, r).items():
            kind = PART_KIND.get(part)
            if not kind:
                continue            # international transfer, Mercury wire fee, refunds: left out
            in_rep = _to_rep(store, r["month"], r.get("fee_ccy"), value)
            if in_rep is None:
                unpriced[r.get("fee_ccy") or "unknown"] += 1
                continue
            fee_rep[(r["month"], r["bank"], kind)] += in_rep
            count[(r["month"], r["bank"], kind)] += 1

    def month_label(m: str) -> str:
        return dt.date.fromisoformat(m + "-01").strftime("%b %y")

    def header(ws, row: int, labels: list[str]) -> None:
        for j, h in enumerate(labels, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[row].height = 32

    wb = Workbook()
    ws = wb.active
    ws.title = "Analysis"
    ws["A1"] = f"Bank and FX fees by month, {rep}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Every fee in the Source Data, from {_start(store)}, split into its parts (the Fees sheet's "
                "breakdown), one bar per bank per month stacked by fee type, then a section per fee type. The "
                "international transfer fee, the Mercury wire fee, unexplained and refunded fees are left out. "
                f"{_who_banks_where()} Fees in another currency are converted at the month's average ECB "
                f"rate (Rates tab), so the {rep} figures are approximate. Built {store['last_refresh'].get('at', '')} UTC.")
    ws["A2"].font = grey
    if unpriced:
        ws["A3"] = "Not converted, no rate for: " + ", ".join(f"{c} ({n})" for c, n in unpriced.items())
        ws["A3"].font = Font(size=9, bold=True, color="C00000")

    # ---- the fee types, in the order and colour the chart and the sections use
    r = 5
    ws.cell(r, 1, "Fee types").font = Font(bold=True, size=12)
    for kind, kb, colour, what, _ in kinds:
        r += 1
        ws.cell(r, 1, kind).font = Font(bold=True, color=colour)
        ws.cell(r, 2, f"{what} ({' and '.join(kb)})")

    # ---- the chart's table: one row per bar
    hdr = (["Bar", "Month", "Bank"]
           + [k if k.startswith(b) else f"{b} {k}" for b, k, _ in series_def] + [f"Total {rep}"])
    first = 4                                  # first series column
    top = r + 2
    header(ws, top, hdr)
    r = top + 1
    for m in months:
        for bank in banks:
            ws.cell(r, 1, f"{month_label(m)[:3]} {bank}")
            ws.cell(r, 2, month_label(m))
            ws.cell(r, 3, bank)
            total = 0.0
            for j, (b, k, _) in enumerate(series_def, start=first):
                v = round(fee_rep.get((m, bank, k), 0.0), 2) if b == bank else None
                ws.cell(r, j, v).number_format = "#,##0"
                total += v or 0
            c = ws.cell(r, len(hdr), round(total, 2))
            c.number_format, c.font = "#,##0", Font(bold=True)
            r += 1
    last = r - 1
    ws.cell(r, 1, "Total").font = Font(bold=True)
    for j in range(first, len(hdr) + 1):
        c = ws.cell(r, j, round(sum(ws.cell(i, j).value or 0 for i in range(top + 1, last + 1)), 2))
        c.number_format, c.font = "#,##0", Font(bold=True)

    chart_row = last + 3
    if months and series_def:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.gapWidth = 60
        chart.title = f"Bank and FX fees by month, {rep} ({', '.join(banks)} bars in that order)"
        chart.y_axis.title = rep
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E6"))
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.legend.position = "b"
        data = Reference(ws, min_col=first, max_col=first - 1 + len(series_def), min_row=top, max_row=last)
        chart.add_data(data, titles_from_data=True)
        # one text label per bar: a two-level axis does not survive Google Sheets
        labels = Reference(ws, min_col=1, min_row=top + 1, max_row=last)
        for series, (_, _, colour) in zip(chart.series, series_def):
            series.cat = AxDataSource(strRef=StrRef(f=str(labels)))
            series.graphicalProperties.solidFill = colour
            series.graphicalProperties.line.solidFill = "FFFFFF"
            series.graphicalProperties.line.width = 12700    # 1pt gap between segments
        chart.height, chart.width = 11, 26
        ws.add_chart(chart, ws.cell(chart_row, 1).coordinate)

    # ---- a section per fee type
    r = chart_row + 24
    for kind, kb, colour, what, _ in kinds:
        kb = tuple(b for b in kb if b in banks)
        c = ws.cell(r, 1, kind)
        c.font = Font(bold=True, size=12, color=colour)
        ws.cell(r, 2, what).font = grey
        r += 1
        header(ws, r, ["Month"] + [h for bank in kb for h in (f"{bank} fees", f"{bank} fees {rep}")])
        r += 1
        start = r
        for m in months + ["Total"]:
            total = m == "Total"
            ws.cell(r, 1, m if total else month_label(m)).font = Font(bold=total)
            j = 2
            for bank in kb:
                for f, fmt in (("rows", "0"), ("fees", "#,##0")):
                    v = (sum(ws.cell(i, j).value or 0 for i in range(start, r)) if total
                         else count.get((m, bank, kind), 0) if f == "rows"
                         else fee_rep.get((m, bank, kind), 0.0))
                    c = ws.cell(r, j, round(v, 2) if isinstance(v, float) else v)
                    c.number_format, c.font = fmt, Font(bold=total)
                    j += 1
            r += 1
        r += 2

    for col, w in ANALYSIS_WIDTHS.items():
        ws.column_dimensions[col].width = w
    from openpyxl.utils import get_column_letter
    for j in range(len(ANALYSIS_WIDTHS) + 1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12

    rates = wb.create_sheet("Rates")
    rates["A1"] = f"{rep} per unit, the month's average ECB reference rate (frankfurter.dev)"
    rates["A1"].font = Font(bold=True)
    ccys = sorted({c for q in (store.get("fx") or {}).values() for c in q})
    for j, h in enumerate(["Month"] + ccys, start=1):
        c = rates.cell(3, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for i, m in enumerate(sorted(store.get("fx") or {}), start=4):
        rates.cell(i, 1, m)
        for j, ccy in enumerate(ccys, start=2):
            rates.cell(i, j, store["fx"][m].get(ccy)).number_format = "0.000000"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def refresh_rates(store: dict) -> str:
    """Re-read the period's monthly rates. One HTTP call; a failure keeps the
    rates already held and says so, never fails the refresh."""
    ccys = {r.get(f) for r in store["rows"].values() for f in ("fee_ccy", "amount_ccy")}
    try:
        # the tariff currencies too, so a fixed price converts into the fee's currency
        ccys |= {str(tariff(b).get("tariff_currency") or "").upper() for b in BANKS}
        store["fx"] = monthly_rates(ccys, _start(store), dt.date.today().isoformat())
        store["fx_base"] = reporting_ccy()
        return "rates refreshed"
    except Exception as e:  # noqa: BLE001
        return f"rates NOT refreshed ({e.__class__.__name__}), the last ones held are used"


def publish(path: Path) -> str:
    from accounting_agent.publish import publish_to_drive
    return publish_to_drive(path, kind="bank_fees") or f"{path.name}: NOT published (no Drive folder configured)"


# ---------------------------------------------------------------- commands

def cmd_refresh(do_publish: bool, start: str | None) -> int:
    store = load_store(start)
    if start:
        store["from"] = start
    changes = refresh(store)
    store["last_refresh"] = {"at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                             "changes": changes}
    rates = refresh_rates(store)
    lines = []
    for path in (build_source(store), build_analysis(store)):
        lines.append(publish(path) if do_publish else f"{path.name}: NOT published (--no-publish)")
    line = "\n".join(lines)
    store["last_refresh"]["published"] = line
    save_store(store)
    print(rates)
    for name, c in changes.items():
        print(f"{name}: {c['added']} added, {c['updated']} updated, {c['removed']} removed, "
              f"API {c['pulled']}")
    print(f"{len(store['rows'])} rows in the store")
    print("Drive: " + line)
    return 0


def cmd_build(do_publish: bool) -> int:
    store = load_store()
    print(refresh_rates(store))
    save_store(store)
    for path in (build_source(store), build_analysis(store)):
        print(path)
        if do_publish:
            print("Drive: " + publish(path))
    return 0


def cmd_status() -> int:
    store = load_store()
    print(f"{len(store['rows'])} rows from {store.get('from') or '-'}, "
          f"last refresh {store['last_refresh'].get('at', 'never')}")
    for key, acc in _accounts().items():
        n = Counter(r["kind"] for r in store["rows"].values()
                    if r["entity"] == acc["short"] and r["bank"] == acc["bank"])
        print(f"{acc['short']} {acc['bank']}: through {store['through'].get(key, '-')}, "
              + ", ".join(f"{k} {v}" for k, v in sorted(n.items())))
    return 0


def summary_line(store: dict | None = None) -> tuple[str, list[str]]:
    """(the report line, manual lines) for the reports run thread."""
    store = store if store is not None else load_store()
    last = store.get("last_refresh") or {}
    at = (last.get("at") or "")[:10]
    manual: list[str] = []
    if not at:
        return "not refreshed yet", ["• bank/fx fees: the store has never been refreshed, run scripts/bank_fees.py refresh on the server"]
    if at != dt.date.today().isoformat():
        manual.append(f"• bank/fx fees: refresh did not run today, data stops at {at}, see data/logs/")
    added = sum(c.get("added", 0) for c in (last.get("changes") or {}).values())
    fees = sum(1 for r in store["rows"].values() if r.get("fee"))
    conv = sum(1 for r in store["rows"].values() if r["kind"] in ("conversion", "foreign payment"))
    through = min(store.get("through", {}).values(), default=at)
    published = last.get("published", "")
    if "NOT published" in published:
        manual.append(f"• bank/fx fees: Drive copy not updated: {published}")
    return (f"source data through {dt.date.fromisoformat(through):%d %b %Y}, {added} new rows · "
            f"{fees:,} fee rows, {conv:,} conversions"), manual


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help", "help"):
        print(__doc__)
        return 0 if argv else 2
    do_publish = "--no-publish" not in argv
    start = argv[argv.index("--from") + 1] if "--from" in argv and argv.index("--from") + 1 < len(argv) else None
    if start:
        dt.date.fromisoformat(start)
    cmd = argv[0]
    if cmd == "refresh":
        return cmd_refresh(do_publish, start)
    if cmd == "build":
        return cmd_build(do_publish)
    if cmd == "status":
        return cmd_status()
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
