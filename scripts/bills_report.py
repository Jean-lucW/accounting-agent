#!/usr/bin/env python
"""The bills payable report: every bill posted in every entity, paid or not,
and for every open bill the reason it is still open, checked against the bank.

    data/reports/Bills Payable.xlsx    the single copy, overwritten
    published to Drive:  <publish folder> / <bills subfolder> /
                         Bills Payable.xlsx   (one copy, overwritten)

WHAT IT ANSWERS. "Which bills have we paid and which have we not?" is a bank
question (docs/BANKING.md): an AUTHORISED bill says nothing about cash. So
every open bill is looked for on every group bank account, and the answer is
one of a fixed set of statuses that say what the reader has to do:

    UNPAID              no movement of that amount on any group bank account
                        since the invoice date: genuinely not paid yet
    MATCH IN XERO       the entity's own bank paid it and the statement line is
                        still unreconciled: match line to bill in Xero
    INTERCOMPANY        another entity's bank paid it: the payer's spend money
                        to the pair's intercompany loan is posted (or not yet),
                        and the bill must be reconciled to intercompany by hand
    DUPLICATE?          the bank movement is already taken by a spend money or
                        by the payment on another bill: one document too many
    VERIFY              several bank lines fit, or a name-only match: look
    NOT VIA BANK        a bill settled by journal, never by the bank: an
                        intercompany recharge, a payroll control bill
    UNVERIFIED          dated before the bank records begin
                        ([company].records_from)

Bank sources by movement date, never mixed (BANKING.md): for each bank account
in config/group.toml, its saved statement CSV up to and including `csv_until`,
the read-only Revolut or Mercury client from the day after (`api_from`). A
`csv` account has no API and is read from its CSV only. A charge in another
currency is recognised at the day's ECB rate. Xero's own SPEND bank
transactions and the payments on paid bills say whether a bank line is already
spoken for; the reconstructed bank feed data/bankfeed/ only says whether the
agent still has the line on its own list.

PAID bills are listed too, with the payment date, the bank account the
payment sits on and whether that payment is reconciled: a paid bill whose
payment is unreconciled is a phantom payment (xero skill rule 7).

READ-ONLY against Xero and the banks. Writes one xlsx on the server and
publishes it to Drive. Runs on the server: it is the only Xero client.

Usage:
    .venv/bin/python scripts/bills_report.py refresh all [--from YYYY-MM-DD] [--no-publish]
    .venv/bin/python scripts/bills_report.py refresh <entity>            # one entity, others from the store
    .venv/bin/python scripts/bills_report.py build [--no-publish]        # rebuild the xlsx from the store, no Xero
    .venv/bin/python scripts/bills_report.py open [<entity>]             # the open bills and their status
    .venv/bin/python scripts/bills_report.py summary                     # counts per entity and status
    .venv/bin/python scripts/bills_report.py notify [<who>] [--dry-run]  # DM the UNPAID list

`notify` sends one Slack DM from the store: the UNPAID bills only, entity by
entity, each as bill date, due date, due status, supplier, invoice number,
amount and currency, and nothing else. With no <who> it goes to
[accounts_payable].unpaid_notify; <who> may be a Slack member ID or a name
from the [slack] tables. It is what closes the scheduled `bills` job so the
payer knows what is to be paid.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from accounting_agent import config  # noqa: E402

OUT_DIR = ROOT / "data" / "reports"
STORE = OUT_DIR / "bills_store.json"
WORKBOOK = OUT_DIR / "Bills Payable.xlsx"
STATEMENTS = ROOT / "data" / "statements"
BANKFEED = ROOT / "data" / "bankfeed"


# Entities, keyed by slug, in config order. The Xero client is always opened
# by the exact organisation name: one name can be a substring of another.
def _ents() -> dict[str, config.Entity]:
    return {e.slug: e for e in config.entities()}


def ORDER() -> list[str]:   # noqa: N802
    return list(_ents())


def SHORT(key: str) -> str:  # noqa: N802
    e = _ents().get(key)
    return e.short if e else key


def _records_from() -> str:
    return config.records_from()


_AP = config.accounts_payable()
# Days a bank movement may precede the invoice date and still be its payment:
# a card is charged on the day, the invoice is issued afterwards (travel
# bookings, hotels).
LEAD_DAYS = int(_AP.get("lead_days", 7))
# The interco funding leg is dated the statement date; allow for a posting lag.
LEG_WINDOW_DAYS = int(_AP.get("leg_window_days", 3))

# Bills that are never settled by a bank payment, matched on the contact or
# the reference.
NOT_VIA_BANK = [(re.compile(n["pattern"], re.I), n.get("reason", "settled by journal, not by a bank payment"))
                for n in (_AP.get("not_via_bank") or []) if n.get("pattern")]
GROUP_NAMES = config.group_name_pattern()
INTERCO_ACCOUNT = re.compile(config.intercompany_settings()["account_name_pattern"], re.I)

# Legal forms (config.LEGAL_FORMS, one list for every script) and the words
# every bank line carries.
STOP = config.LEGAL_FORMS | {"the", "and", "card", "payment", "group"}


def _tokens(s: str) -> set[str]:
    return {t for t in re.split(r"[^a-z0-9]+", (s or "").lower()) if len(t) >= 3 and t not in STOP}


def _xdate(value) -> str:
    """Xero dates arrive as ISO strings or '/Date(ms+0000)/'; return YYYY-MM-DD."""
    if not value:
        return ""
    s = str(value)
    m = re.search(r"/Date\((\d+)", s)
    if m:
        return datetime.fromtimestamp(int(m.group(1)) / 1000, tz=timezone.utc).date().isoformat()
    return s[:10]


def _shift(iso: str, days: int) -> str:
    return (date.fromisoformat(iso) + timedelta(days=days)).isoformat()


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%MZ")


# ------------------------------------------------------------------ Xero pull

def _client(key: str):
    from accounting_agent.xero import XeroClient
    return XeroClient(_ents()[key].xero_name)


def pull_entity(key: str, date_from: str) -> dict:
    """Every ACCPAY bill AUTHORISED (any date) or PAID (dated >= date_from),
    the payments on them, the entity's SPEND bank transactions since the
    oldest open bill, the chart and the org short code."""
    from accounting_agent.xero.banking import get_bank_transactions, get_payments
    from accounting_agent.xero.purchases import get_bills
    c = _client(key)
    org = (c.get("Organisation").get("Organisations") or [{}])[0]
    chart = c.get("Accounts").get("Accounts", [])
    names = {a.get("Code"): a.get("Name") for a in chart if a.get("Code")}
    acct_by_id = {a.get("AccountID"): a for a in chart}
    bank_names = {a["AccountID"]: f"{a.get('Name')}" for a in chart if a.get("Type") == "BANK"}

    open_bills = get_bills(c, status="AUTHORISED", unpaid_only=True)
    paid_bills = get_bills(c, status="PAID", date_from=date_from)
    payments = get_payments(c, from_date=date_from)
    by_invoice: dict[str, list[dict]] = defaultdict(list)
    for p in payments:
        if (p.get("Status") or "AUTHORISED") != "AUTHORISED":
            continue
        inv = (p.get("Invoice") or {}).get("InvoiceID")
        if not inv:
            continue
        acct = p.get("Account") or {}
        a = acct_by_id.get(acct.get("AccountID")) or {}
        by_invoice[inv].append({
            "date": _xdate(p.get("Date")),
            "amount": float(p.get("Amount") or 0),
            "account": (a.get("Name") or names.get(acct.get("Code")) or acct.get("Code") or "").replace("••", "xx"),
            "account_code": a.get("Code") or acct.get("Code") or "",
            "is_bank": a.get("Type") == "BANK",
            "reconciled": bool(p.get("IsReconciled")),
            "reference": p.get("Reference") or "",
        })

    oldest_open = min([_xdate(b.get("Date")) for b in open_bills] or [date_from])
    spend_from = min(_shift(oldest_open, -LEAD_DAYS), date_from)
    spends = get_bank_transactions(c, from_date=spend_from, where='Type=="SPEND"', status="AUTHORISED")
    spend_rows = []
    for t in spends:
        ba = t.get("BankAccount") or {}
        codes = []
        for li in t.get("LineItems") or []:
            code = li.get("AccountCode")
            if code and code not in codes:
                codes.append(code)
        spend_rows.append({
            "date": _xdate(t.get("Date")),
            "amount": float(t.get("Total") or 0),
            "currency": t.get("CurrencyCode") or "",
            "bank_account": bank_names.get(ba.get("AccountID")) or ba.get("Name") or ba.get("Code") or "",
            "contact": ((t.get("Contact") or {}).get("Name") or ""),
            "reference": t.get("Reference") or "",
            "accounts": [names.get(code, code) for code in codes],
            "reconciled": bool(t.get("IsReconciled")),
        })

    def bill_row(b: dict, status: str) -> dict:
        codes, descs = [], []
        for li in b.get("LineItems") or []:
            code = li.get("AccountCode")
            if code and code not in codes:
                codes.append(code)
            d = " ".join((li.get("Description") or "").split())
            if d and d not in descs:
                descs.append(d)
        iid = b.get("InvoiceID") or ""
        return {
            "id": iid,
            "status": status,
            "date": _xdate(b.get("Date")),
            "due": _xdate(b.get("DueDate")),
            "supplier": " ".join(((b.get("Contact") or {}).get("Name") or "").split()),
            "number": b.get("InvoiceNumber") or "",
            "reference": b.get("Reference") or "",
            "currency": b.get("CurrencyCode") or "",
            "total": float(b.get("Total") or 0),
            "paid": float(b.get("AmountPaid") or 0),
            "credited": float(b.get("AmountCredited") or 0),
            "due_amount": float(b.get("AmountDue") or 0),
            "accounts": [names.get(code, code) for code in codes],
            "account_codes": codes,
            "description": " | ".join(descs)[:500],
            "attachment": bool(b.get("HasAttachments")),
            "updated": _xdate(b.get("UpdatedDateUTC")),
            "fully_paid_on": _xdate(b.get("FullyPaidOnDate")),
            "payments": by_invoice.get(iid, []),
        }

    bills = [bill_row(b, "OPEN") for b in open_bills] + [bill_row(b, "PAID") for b in paid_bills]
    return {
        "pulled": _now(),
        "from": date_from,
        "short_code": org.get("ShortCode") or "",
        "base_currency": org.get("BaseCurrency") or "",
        "bills": bills,
        "spends": spend_rows,
    }


# ------------------------------------------------------------------ the bank

def _num(v) -> float:
    try:
        return float(str(v or 0).replace(",", ""))
    except ValueError:
        return 0.0


def _csv_lines(key: str, bank: config.BankAccount) -> list[dict]:
    """Settled movements from the account's saved statement CSV, up to and
    including its csv_until (every row for a csv-only account without one).
    The schema is recognised from the header: Revolut, Mercury, or the generic
    one docs/BANKING.md defines for a bank with no API."""
    if not bank.statement_csv:
        return []
    if bank.has_api and not bank.csv_until:
        return []   # API only: the CSV is not the record for any date
    path = STATEMENTS / bank.statement_csv
    if not path.exists():
        return []
    out = []
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if rows and "Date completed (UTC)" in rows[0]:            # Revolut
        for r in rows:
            if (r.get("State") or "").upper() != "COMPLETED":
                continue
            try:
                amt = float(r.get("Amount") or 0)
            except ValueError:
                continue
            if not amt:
                continue
            ccy = r.get("Payment currency") or ""
            if bank.currency and ccy.upper() != bank.currency.upper():
                continue
            oc, oa = r.get("Orig currency") or "", r.get("Orig amount") or ""
            out.append({
                "entity": key, "date": (r.get("Date completed (UTC)") or r.get("Date started (UTC)") or "")[:10],
                "payee": " ".join((r.get("Description") or "").split()),
                "amount": amt, "currency": ccy,
                "orig_amount": float(oa) if oa else None, "orig_currency": oc,
                "account": r.get("Account") or "", "who": r.get("Payer") or r.get("Card label") or "",
                "reference": r.get("Reference") or "", "source": "statement csv",
            })
    elif rows and "Date (UTC)" in rows[0]:                      # Mercury
        for r in rows:
            if (r.get("Status") or "").lower() != "sent":
                continue
            try:
                amt = float(r.get("Amount") or 0)
            except ValueError:
                continue
            if not amt:
                continue
            d = r.get("Date (UTC)") or ""
            m = re.match(r"(\d{2})-(\d{2})-(\d{4})", d)
            iso = f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else d[:10]
            out.append({
                "entity": key, "date": iso,
                "payee": " ".join((r.get("Description") or r.get("Bank Description") or "").split()),
                "amount": amt, "currency": bank.currency or "USD", "orig_amount": None,
                "orig_currency": r.get("Original Currency") or "",
                "account": r.get("Source Account") or "", "who": r.get("Name On Card") or "",
                "reference": " ".join(filter(None, (r.get("Reference"), r.get("Note"), r.get("Bank Description")))),
                "source": "statement csv",
            })
    else:                                                       # generic (docs/BANKING.md)
        for r in rows:
            low = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
            if (low.get("status") or "").lower() in ("pending", "failed", "cancelled", "declined", "reverted"):
                continue
            amt = _num(low.get("amount"))
            if not amt:
                continue
            ccy = (low.get("currency") or bank.currency or _ents()[key].base_currency).upper()
            oa = low.get("original amount") or ""
            out.append({
                "entity": key, "date": (low.get("date") or "")[:10],
                "payee": " ".join((low.get("description") or low.get("payee") or "").split()),
                "amount": amt, "currency": ccy,
                "orig_amount": _num(oa) if oa else None,
                "orig_currency": (low.get("original currency") or "").upper(),
                "account": low.get("account") or bank.label, "who": low.get("payer") or low.get("card") or "",
                "reference": low.get("reference") or "", "source": "statement csv",
            })
    if bank.csv_until:
        out = [l for l in out if l["date"] <= bank.csv_until]
    return out


def _api_lines(key: str, bank: config.BankAccount) -> list[dict]:
    """Settled movements from the bank API, from the account's api_from
    (or the start of the records when it has no CSV cut-over)."""
    if not bank.has_api:
        return []
    source = bank.provider
    if source == "revolut":
        from accounting_agent.revolut import RevolutClient
        client = RevolutClient(bank.slug)
        cash = {"completed"}
    else:
        from accounting_agent.mercury import MercuryClient
        client = MercuryClient(bank.slug)
        cash = {"sent"}
    start = bank.csv_until or _records_from()
    out = []
    for l in client.statement_lines(from_date=start):
        if (l.get("state") or "").lower() not in cash:
            continue
        amt = float(l.get("amount") or 0)
        if not amt:
            continue
        d = l.get("date") or ""
        if (bank.csv_until and d <= bank.csv_until) or (not bank.csv_until and d < start):
            continue
        if bank.currency and (l.get("currency") or "").upper() != bank.currency.upper():
            continue
        acct = bank.account_heading(l)
        out.append({
            "entity": key, "date": d,
            "payee": " ".join((l.get("merchant") or l.get("description") or l.get("counterparty") or "").split()),
            "amount": amt, "currency": l.get("currency") or "",
            "orig_amount": float(l["bill_amount"]) if l.get("bill_amount") else None,
            "orig_currency": l.get("bill_currency") or "",
            "account": acct, "who": l.get("cardholder") or "",
            "reference": " ".join(filter(None, (l.get("reference"), l.get("description")))),
            "source": f"{source} api",
        })
    return out


def bank_history() -> tuple[list[dict], dict[str, str]]:
    """Every settled movement on every group account, and per entity the date
    the record runs to."""
    lines, upto = [], {}
    for key, ent in _ents().items():
        if not ent.bank_accounts:
            continue
        reached = []
        for bank in ent.bank_accounts:
            lines.extend(_csv_lines(key, bank))
            if bank.has_api:
                try:
                    lines.extend(_api_lines(key, bank))
                    reached.append(date.today().isoformat())
                except Exception as e:   # noqa: BLE001 - the bank being down must not lose the report
                    reached.append(f"{bank.csv_until or 'nothing'} ({bank.label} api failed: {e.__class__.__name__})")
            else:
                reached.append(bank.csv_until or f"end of {bank.statement_csv or 'no statement'}")
        upto[key] = "; ".join(reached)
    return lines, upto


def bank_feed_lines() -> list[dict]:
    """What is still unreconciled, per data/bankfeed."""
    out = []
    head = re.compile(r"^## (.+?)(?: · last checked (\S+))?\s*$")
    if not BANKFEED.exists():
        return out
    for path in sorted(BANKFEED.glob("*.md")):
        account = ""
        for raw in path.read_text().splitlines():
            m = head.match(raw)
            if m:
                account = m.group(1)
                continue
            parts = [p.strip() for p in raw.split("|")]
            if len(parts) < 4 or not re.match(r"\d{4}-\d{2}-\d{2}", parts[0]):
                continue
            try:
                ccy, amt = parts[3].split(" ", 1)
                out.append({"entity": path.stem, "account": account, "date": parts[0], "payee": parts[1],
                            "direction": parts[2], "amount": float(amt.replace(",", "")), "currency": ccy,
                            "who": parts[4] if len(parts) > 4 else "", "note": parts[6] if len(parts) > 6 else ""})
            except ValueError:
                continue
    return out


# ------------------------------------------------------------ classification
#
# The bank feed (data/bankfeed) is the AGENT's to-do list, not Xero's
# reconcile screen: a line leaves it the moment the run posts the bill for it,
# before the user has matched anything. So "not in the feed" never means
# "reconciled in Xero". What does say a bank line is spoken for is Xero
# itself: a spend money of that amount and date, or a payment on another bill.

# Words that two unrelated names share all the time. A match on these alone
# is not a name match. The built-in list is the shared legal forms
# (config.LEGAL_FORMS) plus neutral company and place words; a group adds the
# words common in its own supplier base through [accounts_payable]
# generic_supplier_words in config/group.toml.
GENERIC_BUILTIN = config.LEGAL_FORMS | frozenset({
    "the", "and", "group", "holdings",
    "services", "service", "solutions", "systems", "system", "international", "global",
    "uk", "us", "usa", "eu", "europe", "management", "partners", "consulting", "office",
    "online", "digital", "technologies", "technology", "network", "operations", "payments",
    "bank", "store", "shop", "city", "centre", "center", "street", "square", "travel",
    "hotel", "hotels", "taxi", "transport", "airlines", "airways", "air",
})


def _generic_words() -> frozenset[str]:
    return GENERIC_BUILTIN | config.legal_form_words()


GENERIC = _generic_words()
# How far a bank movement in another currency may sit from the ECB rate and
# still be the same charge: card schemes and banks add a spread. The wider
# band (a tip on a card, a scheme rate on a bad day) needs the payee to read as
# the supplier as well.
FX_TOLERANCE = 0.06
FX_TOLERANCE_NAMED = 0.12
FEE_LINE = re.compile(r"\bfee\b", re.I)


def _shared(a: set[str], b: set[str]) -> set[str]:
    """Words the two names share: equal, or one a 5+ letter prefix of the other
    (the bank truncates: "Transporta" for "Transportation")."""
    out = set()
    for x in a:
        for y in b:
            if x == y or (len(x) >= 5 and len(y) >= 5 and (x.startswith(y) or y.startswith(x))):
                out.add(x)
    return out


def _name_score(payee: str, supplier: str) -> int:
    """2 per distinctive shared word, 1 per generic one. 0 = amount only."""
    shared = _shared(_tokens(payee), _tokens(supplier))
    return 2 * len(shared - GENERIC) + len(shared & GENERIC)


def _fits(line: dict, ccy: str, amount: float) -> bool:
    """The movement is this amount in this currency: on the account side, or
    as the original charge when the account is in another currency. A GBP
    account-side amount that the bank says was a EUR charge is not a GBP bill."""
    oa, oc = line.get("orig_amount"), line.get("orig_currency") or ""
    if oa is not None and oc and oc == ccy and abs(abs(oa) - amount) < 0.005:
        return True
    if abs(abs(line["amount"]) - amount) < 0.005 and line["currency"] == ccy:
        return not (oa is not None and oc and oc != ccy)
    return False


_FX_CACHE: dict = {}


def _fx(base: str, quote: str, day: str) -> float | None:
    """How many `quote` one `base` bought on `day` (ECB via frankfurter).
    None when the rate cannot be had."""
    if base == quote:
        return 1.0
    key = f"{base}/{quote}/{day}"
    if key in _FX_CACHE:
        return _FX_CACHE[key]
    rate = None
    try:
        import urllib.parse
        import urllib.request
        url = f"https://api.frankfurter.dev/v1/{day}?" + urllib.parse.urlencode({"base": base, "symbols": quote})
        req = urllib.request.Request(url, headers={"Accept": "application/json",
                                                   "User-Agent": "accounting-agent-bills-report/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            rate = (json.load(resp).get("rates") or {}).get(quote)
    except Exception:   # noqa: BLE001 - no rate means no FX match, never a failed report
        rate = None
    _FX_CACHE[key] = rate
    return rate


def _fx_fits(line: dict, ccy: str, amount: float, tolerance: float = FX_TOLERANCE) -> bool:
    """A payment in another currency: what the bank charged (the original
    amount when it says so, else the account side) agrees with the bill at
    the day's rate within the tolerance."""
    if FEE_LINE.search(line["payee"] or ""):
        return False
    lc, la = _line_amount(line)
    if lc == ccy:
        return False
    rate = _fx(ccy, lc, line["date"])
    if not rate:
        return False
    expected = amount * rate
    return abs(la - expected) <= max(tolerance * expected, 0.05)


def _in_feed(feed: list[dict], line: dict) -> bool:
    return any(f["entity"] == line["entity"] and f["date"] == line["date"] and f["direction"] == "spent"
               and f["currency"] == line["currency"] and abs(f["amount"] - abs(line["amount"])) < 0.005
               for f in feed)


def _evidence(line: dict) -> str:
    who = f" · {line['who']}" if line.get("who") else ""
    orig = (f" (charged {line['orig_currency']} {abs(line['orig_amount']):,.2f})"
            if line.get("orig_amount") and line.get("orig_currency") != line["currency"] else "")
    return (f"{SHORT(line['entity'])} · {line['account']} · {line['date']} · {line['payee']} · "
            f"{line['currency']} {abs(line['amount']):,.2f}{orig}{who} · {line['source']}")


def _spend_near(spends: list[dict], line: dict, interco: bool | None = None) -> dict | None:
    """A spend money in the paying entity of the line's amount within the leg
    window; interco=True restricts to intercompany-coded, False to anything else."""
    lo, hi = _shift(line["date"], -LEG_WINDOW_DAYS), _shift(line["date"], LEG_WINDOW_DAYS)
    for s in spends:
        if not (lo <= s["date"] <= hi) or abs(s["amount"] - abs(line["amount"])) > 0.005:
            continue
        is_ic = any(INTERCO_ACCOUNT.search(a or "") for a in s["accounts"])
        if interco is None or is_ic == interco:
            return s
    return None


def _payment_near(paid_bills: list[dict], line: dict, exclude_id: str) -> dict | None:
    """Another bill whose payment took this bank line."""
    lo, hi = _shift(line["date"], -LEG_WINDOW_DAYS), _shift(line["date"], LEG_WINDOW_DAYS)
    for b in paid_bills:
        if b["id"] == exclude_id:
            continue
        for p in b.get("payments") or []:
            if p["is_bank"] and lo <= p["date"] <= hi and abs(p["amount"] - abs(line["amount"])) < 0.005:
                return b
    return None


def _batches(open_by_entity: dict[str, list[dict]], history: list[dict]) -> dict[str, tuple[dict, list[str]]]:
    """One bank movement paying several bills of one supplier: bill id ->
    (line, the invoice numbers paid together)."""
    from itertools import combinations
    out: dict[str, tuple[dict, list[str]]] = {}
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for key, bills in open_by_entity.items():
        for b in bills:
            groups[(key, " ".join(_tokens(b["supplier"]) - GENERIC) or b["supplier"].lower(), b["currency"])].append(b)
    for (key, _, ccy), bills in groups.items():
        if len(bills) < 2 or len(bills) > 12:
            continue
        floor = _shift(min(b["date"] for b in bills), -LEAD_DAYS)
        sums = {}
        for k in range(2, min(5, len(bills)) + 1):
            for combo in combinations(bills, k):
                if any(c["id"] in out for c in combo):
                    continue
                sums.setdefault(round(sum(c["due_amount"] for c in combo), 2), combo)
        for line in history:
            if line["amount"] >= 0 or line["date"] < floor:
                continue
            for total, combo in sums.items():
                if _fits(line, ccy, total) and not any(c["id"] in out for c in combo):
                    numbers = [c["number"] or c["supplier"] for c in combo]
                    for c in combo:
                        out[c["id"]] = (line, numbers)
                    break
    return out


def _line_key(line: dict) -> tuple:
    return (line["entity"], line["account"], line["date"], line["payee"], round(line["amount"], 2))


def _alnum(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def _number_on(line: dict, number: str) -> bool:
    """The invoice number is written on the bank line (payee or reference).
    Short or purely numeric tokens under 5 digits are too easy to hit."""
    n = _alnum(number)
    if len(n) < 5 or (n.isdigit() and len(n) < 6):
        return False
    return n in _alnum(line.get("payee", "")) or n in _alnum(line.get("reference", ""))


# A transfer to a named person, an exchange between the entity's own pockets
# or a receipt is never a supplier's card charge.
NOT_CARD = re.compile(r"^(to |payment from |exchange)|→|->", re.I)


def _card_like(line: dict) -> bool:
    return not NOT_CARD.search(line.get("payee") or "")


def _in_bill_ccy(line: dict, ccy: str) -> float | None:
    """What the bank charged, expressed in the bill's currency at the day's rate."""
    lc, la = _line_amount(line)
    if lc == ccy:
        return la
    rate = _fx(lc, ccy, line["date"])
    return la * rate if rate else None


def _line_amount(line: dict) -> tuple[str, float]:
    """What was charged: the original currency and amount when the bank says
    so, else the account side."""
    if line.get("orig_amount") is not None and line.get("orig_currency"):
        return line["orig_currency"], abs(line["orig_amount"])
    return line["currency"], abs(line["amount"])


def classify(key: str, bill: dict, history: list[dict], feed: list[dict], store: dict,
             upto: dict[str, str], batch: tuple[dict, list[str]] | None,
             claimed: set[tuple] | None = None) -> tuple[str, str, str]:
    """-> (status, reason and action, bank evidence) for an OPEN bill.
    `claimed` holds the bank lines earlier bills already resolved to; a
    resolved line is added to it so two bills never share one movement.

    Evidence of payment, broadest first: the invoice number on a bank line;
    the amount, in the bill's currency or converted at the day's rate; the
    supplier's name on a bank line at some other amount."""
    claimed = claimed if claimed is not None else set()
    for rx, why in NOT_VIA_BANK:
        if rx.search(bill["supplier"]) or rx.search(bill.get("reference") or ""):
            return "NOT VIA BANK", why, ""
    if GROUP_NAMES.search(bill["supplier"]):
        return "NOT VIA BANK", "intercompany recharge: settled through the intercompany account by journal, not the bank", ""
    records_from = _records_from()
    if bill["date"] < records_from:
        return "UNVERIFIED", f"before the bank records ({records_from})", ""

    ents = store["entities"]
    spends = {k: v.get("spends", []) for k, v in ents.items()}
    paid = {k: [b for b in v.get("bills", []) if b["status"] == "PAID"] for k, v in ents.items()}
    floor = _shift(bill["date"], -LEAD_DAYS)
    near_hi = _shift(bill["date"], LEAD_DAYS)
    ccy, due = bill["currency"], round(bill["due_amount"], 2)
    amounts = {due} | ({round(bill["total"], 2)} if (bill["paid"] or bill["credited"]) else set())
    partial = f" Part-paid: {ccy} {bill['paid'] + bill['credited']:,.2f} settled, {ccy} {due:,.2f} due." if (bill["paid"] or bill["credited"]) else ""
    d = lambda iso: f"{int(iso[8:10])} {date.fromisoformat(iso).strftime('%b')}"   # noqa: E731

    def spoken_for(h: dict) -> bool:
        """Taken in its own entity by another bill's payment or a spend money."""
        return bool(_payment_near(paid.get(h["entity"], []), h, bill["id"])
                    or _spend_near(spends.get(h["entity"], []), h))

    def usual_account() -> str:
        sup = _tokens(bill["supplier"]) - GENERIC
        last = None
        for b in paid.get(key, []):
            if not sup or not (_shared(_tokens(b["supplier"]), sup) - GENERIC):
                continue
            for pmt in b.get("payments") or []:
                if pmt["account"] and (last is None or pmt["date"] > last["date"]):
                    last = pmt
        return f" Last paid from {last['account']}, {d(last['date'])}." if last else ""

    def resolve(h: dict, note: str) -> tuple[str, str, str]:
        """One bank movement is this bill's payment; say what state it is in."""
        claimed.add(_line_key(h))
        ev = _evidence(h)
        score = _name_score(h["payee"], bill["supplier"])
        payee_note = "" if score or h.get("_number") else f" Payee {h['payee']}: confirm."
        if h["entity"] == key:
            other = _payment_near(paid[key], h, bill["id"])
            if other:
                return ("DUPLICATE?", f"bank line {d(h['date'])} already matched to {other['supplier']} bill "
                                      f"{other['number']}: duplicate, void one.{note}{partial}", ev)
            coded = _spend_near(spends[key], h)
            if coded:
                return ("DUPLICATE?", f"bank line {d(h['date'])} already coded as spend money to "
                                      f"{', '.join(coded['accounts']) or '?'}: duplicate, void one or recode.{note}{partial}", ev)
            return ("MATCH IN XERO", f"paid {d(h['date'])} from own bank: match the line to this bill.{note}{payee_note}{partial}", ev)
        payer = SHORT(h["entity"])
        leg = _spend_near(spends.get(h["entity"], []), h, interco=True)
        if score == 0 and not h.get("_number") and not leg:
            return ("VERIFY", f"{payer} paid the amount {d(h['date'])} to {h['payee']}: same bill? confirm.{note}{partial}", ev)
        if leg:
            state = "reconciled" if leg["reconciled"] else "posted"
            return ("INTERCOMPANY", f"paid by {payer} {d(h['date'])}, intercompany loan leg {state}: reconcile bill to "
                                    f"intercompany.{note}{payee_note}{partial}", ev)
        other = _payment_near(paid.get(h["entity"], []), h, bill["id"])
        if other:
            return ("DUPLICATE?", f"paid by {payer} {d(h['date'])}, which holds its own paid bill {other['number']} "
                                  f"({other['supplier']}): one bill is in the wrong entity, void it.{note}{partial}", ev)
        coded = _spend_near(spends.get(h["entity"], []), h, interco=False)
        if coded:
            return ("INTERCOMPANY", f"paid by {payer} {d(h['date'])}, coded there to {', '.join(coded['accounts']) or '?'}: "
                                    f"recode to the intercompany loan, then reconcile bill to intercompany.{note}{partial}", ev)
        return ("INTERCOMPANY", f"paid by {payer} {d(h['date'])}, no intercompany loan leg yet (bill-payments posts it): "
                                f"then reconcile bill to intercompany.{note}{partial}", ev)

    if batch:
        line, numbers = batch
        return resolve(line, f" One payment of {line['currency']} {abs(line['amount']):,.2f} covers {', '.join(numbers)}.")

    # ---- gather every kind of evidence
    cands = []
    for l in history:
        if l["amount"] >= 0 or l["date"] < floor or _line_key(l) in claimed:
            continue
        tags = set()
        if any(_fits(l, ccy, a) for a in amounts):
            tags.add("amount")
        elif l["date"] <= near_hi and _fx_fits(l, ccy, due):
            tags.add("fx")
        score = _name_score(l["payee"], bill["supplier"])
        if l["date"] <= near_hi and bill["number"] and _number_on(l, bill["number"]):
            tags.add("number")
        if l["date"] <= near_hi and score >= 2:
            tags.add("name")
            if "fx" not in tags and "amount" not in tags and _fx_fits(l, ccy, due, FX_TOLERANCE_NAMED):
                tags.add("fx")
        if "fx" in tags and score == 0 and "number" not in tags:
            # a converted amount with a stranger's name only counts for a card
            # charge on this entity's own account that nothing else has taken
            if l["entity"] != key or not _card_like(l) or spoken_for(l):
                tags.discard("fx")
        if not tags:
            continue
        l["_score"], l["_tags"], l["_number"] = score, tags, "number" in tags
        cands.append(l)

    # a movement of the amount in another entity, payee unlike the supplier and
    # already taken there by its own bill or spend money, is a coincidence
    aside = [c for c in cands if c["entity"] != key and c["_score"] == 0 and not c["_number"]
             and spoken_for(c) and not _spend_near(spends.get(c["entity"], []), c, interco=True)]
    cands = [c for c in cands if c not in aside]

    def pick(rows: list[dict]) -> dict | None:
        """The one movement that is clearly this bill, or None if it is a toss-up."""
        if len(rows) == 1:
            return rows[0]
        same = [r for r in rows if r["entity"] == key]
        if len(same) == 1 and all(r["_score"] == 0 for r in rows if r is not same[0]):
            return same[0]
        rows = sorted(rows, key=lambda r: (-r["_score"], r.get("_dev", 0), r["date"]))
        a, b = rows[0], rows[1]
        if a["_score"] >= 1 and a["_score"] > b["_score"]:
            return a
        if a["_score"] == b["_score"] and a.get("_dev", 0) <= FX_TOLERANCE < b.get("_dev", 0):
            return a
        return None

    numbered = [c for c in cands if c["_number"]]
    if numbered:
        h = pick(numbered) or numbered[0]
        lc, la = _line_amount(h)
        note = "" if "amount" in h["_tags"] else f" Invoice number on the bank line; amount there {lc} {la:,.2f}."
        return resolve(h, note or " Invoice number on the bank line.")

    exact = [c for c in cands if "amount" in c["_tags"]]
    if exact:
        h = pick(exact)
        if h:
            return resolve(h, "")
        return ("VERIFY", f"{len(exact)} bank lines fit the amount: pick the one.{partial}",
                "; ".join(_evidence(c) for c in exact[:5]))

    fx = [c for c in cands if "fx" in c["_tags"]]
    for c in fx:
        lc, la = _line_amount(c)
        rate = _fx(ccy, lc, c["date"]) or 0
        c["_dev"] = abs(la - due * rate) / (due * rate) if rate else 1.0
    if fx:
        h = pick(fx)
        if h and h["_score"] == 0:
            # the amount converts, the name says nothing: a question, not a match
            lc, la = _line_amount(h)
            return ("VERIFY", f"1 bank line fits at the day's rate: {h['payee']} {lc} {la:,.2f} {d(h['date'])}. "
                              f"Same bill?{partial}", _evidence(h))
        if h:
            lc, la = _line_amount(h)
            return resolve(h, f" Charged as {lc} {la:,.2f}.")
        return ("VERIFY", f"{len(fx)} bank lines fit the amount at the day's rate: pick the one.{partial}",
                "; ".join(_evidence(c) for c in fx[:5]))

    # the supplier's name at another amount: only an amount in the bill's
    # range says anything (a batch or a tip is more, never a tenth)
    named = []
    for c in cands:
        if "name" not in c["_tags"] or spoken_for(c):
            continue
        v = _in_bill_ccy(c, ccy)
        if v is not None and 0.9 * due <= v <= 5 * due:
            named.append(c)
    if named:
        named.sort(key=lambda c: c["date"])
        shown = ", ".join(f"{_line_amount(c)[0]} {_line_amount(c)[1]:,.2f} {d(c['date'])}" for c in named[:3])
        return ("VERIFY", f"supplier on the bank at other amounts ({shown}), not matched to any bill: batch, "
                          f"tip or another invoice?{partial}", "; ".join(_evidence(c) for c in named[:4]))

    aside_note = ""
    if aside:
        a = aside[0]
        aside_note = f" ({SHORT(a['entity'])} {a['currency']} {abs(a['amount']):,.2f} to {a['payee']} {d(a['date'])} is unrelated.)"
    return ("UNPAID", f"no payment found since {d(floor)}.{usual_account()}{aside_note}{partial}", "")


def paid_detail(bill: dict) -> tuple[str, str, str, bool]:
    """-> (paid on, paid from, payment state, reconciled) for a PAID bill.
    Reconciled means every bank payment on it is matched to a statement line;
    a bill settled from a loan or control account has no bank line to match."""
    pays = bill.get("payments") or []
    if not pays:
        if bill["credited"] and bill["credited"] >= bill["total"] - 0.005:
            return bill["fully_paid_on"], "credit note", "settled by credit note", True
        return bill["fully_paid_on"], "", "no payment record in the window", True
    dates = sorted(p["date"] for p in pays)
    accounts = []
    for p in pays:
        if p["account"] and p["account"] not in accounts:
            accounts.append(p["account"])
    unrec = [p for p in pays if p["is_bank"] and not p["reconciled"]]
    if unrec:
        shown = "; ".join(f"{p['date']} {p['amount']:,.2f}" for p in unrec)
        return dates[-1], ", ".join(accounts), f"PHANTOM? bank payment not reconciled to a statement line: {shown}", False
    if any(p["is_bank"] for p in pays):
        return dates[-1], ", ".join(accounts), "payment reconciled to the bank", True
    return dates[-1], ", ".join(accounts), "settled from a loan or control account", True


# ------------------------------------------------------------------- store

def load_store() -> dict:
    if STORE.exists():
        return json.loads(STORE.read_text())
    return {"entities": {}}


def save_store(store: dict) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    STORE.write_text(json.dumps(store, indent=1, sort_keys=True))
    STORE.chmod(0o600)


def assess(store: dict) -> dict:
    """Classify every open bill against the bank and record it on the store."""
    history, upto = bank_history()
    feed = bank_feed_lines()
    open_by_entity = {k: [b for b in v["bills"] if b["status"] == "OPEN"] for k, v in store["entities"].items()}
    batches = _batches(open_by_entity, history)
    counts: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    claimed: set[tuple] = set()
    for key, ent in store["entities"].items():
        for b in sorted(ent["bills"], key=lambda b: (b["date"], b["supplier"])):
            if b["status"] == "OPEN":
                st, why, ev = classify(key, b, history, feed, store, upto, batches.get(b["id"]), claimed)
                b["state"], b["why"], b["evidence"] = st, why, ev
                counts[key][st] += 1
                b["reconciled"] = False
            else:
                on, frm, state, rec = paid_detail(b)
                b["state"], b["why"], b["evidence"] = "PAID", state, ""
                b["paid_on"], b["paid_from"], b["reconciled"] = on, frm, rec
                counts[key]["PAID"] += 1
                if not rec:
                    counts[key]["PAID, UNRECONCILED"] += 1
    store["assessed"] = _now()
    store["bank_upto"] = upto
    store["bank_lines"] = len(history)
    store["feed_lines"] = len(feed)
    store["counts"] = {k: dict(v) for k, v in counts.items()}
    return store


# ---------------------------------------------------------------- workbook

STATE_ORDER = ["INTERCOMPANY", "DUPLICATE?", "VERIFY", "MATCH IN XERO", "UNPAID",
               "NOT VIA BANK", "UNVERIFIED", "PAID"]
STATE_HELP = {
    "UNPAID": "no movement of the amount on any group bank account since the invoice date: genuinely unpaid",
    "MATCH IN XERO": "the entity's own bank paid it; the statement line is unreconciled: match it to the bill in Xero",
    "INTERCOMPANY": "another entity's bank paid it: reconcile the bill to intercompany by hand (the payer's intercompany loan leg is posted by the bill-payments check)",
    "DUPLICATE?": "the bank line is already taken by a spend money or by the payment on another bill: one document too many, void one in a bookkeeping run",
    "VERIFY": "several bank lines fit and none reads as the supplier more than the others, or a movement of the amount in another entity has a payee unlike the supplier, or the supplier appears on the bank at other amounts not matched to any bill: look before acting",
    "NOT VIA BANK": "settled by journal, never by a bank payment: payroll control bills, intercompany recharges",
    "UNVERIFIED": "dated before the bank records begin",
    "PAID": "Xero shows it paid; hidden by the Reconciled filter unless the bank payment is not reconciled to a statement line (PHANTOM?)",
}


def _cutover_note() -> str:
    """How each bank account is read, for the Summary tab."""
    parts = []
    for e in config.entities():
        for b in e.bank_accounts:
            if b.has_api and b.csv_until:
                parts.append(f"{e.short} {b.label}: CSV to {b.csv_until}, API after")
            elif b.has_api:
                parts.append(f"{e.short} {b.label}: API")
            else:
                parts.append(f"{e.short} {b.label}: CSV only")
    return "; ".join(parts) or "no bank accounts configured"


def build_workbook(store: dict, out: Path = WORKBOOK) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.filters import FilterColumn, Filters

    HEAD = PatternFill("solid", fgColor="203864")
    FILLS = {
        "UNPAID": PatternFill("solid", fgColor="FFF2CC"),
        "MATCH IN XERO": PatternFill("solid", fgColor="DDEBF7"),
        "INTERCOMPANY": PatternFill("solid", fgColor="F8CBAD"),
        "DUPLICATE?": PatternFill("solid", fgColor="F4B6B6"),
        "VERIFY": PatternFill("solid", fgColor="FCE4D6"),
        "NOT VIA BANK": PatternFill("solid", fgColor="E2EFDA"),
        "UNVERIFIED": PatternFill("solid", fgColor="D9D9D9"),
        "PAID": PatternFill("solid", fgColor="C6E0B4"),
    }
    today = date.today().isoformat()
    ents = _ents()
    order = ORDER()
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{config.company_name()} bills payable: every bill posted, paid or open, and why an open bill is still open"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Built {store.get('assessed', '')} UTC. Every open bill searched for on every group bank account "
                f"({_cutover_note()}): amount in any currency, invoice number, supplier name. "
                f"Entity tabs show unreconciled bills only; clear the Reconciled filter to see the paid ones.")
    no_bank = [e.short for e in ents.values() if not e.bank_accounts]
    ws["A3"] = ("Bank read to " + ", ".join(f"{SHORT(k)} {v}" for k, v in (store.get("bank_upto") or {}).items())
                + f"; {store.get('bank_lines', 0):,} movements."
                + (f" No bank account: {', '.join(no_bank)} (checked against every group account)." if no_bank else "")
                + " Read-only, nothing changed in Xero.")
    for r in (2, 3):
        ws[f"A{r}"].font = Font(size=9, italic=True, color="555555")

    hdr = ["Entity", "Open bills", *STATE_ORDER[:-1], "Overdue", "Paid, unreconciled", "Open by currency", "Xero pulled"]
    r = 5
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(r, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 32
    r += 1
    tot = defaultdict(int)
    for key in order:
        ent = store["entities"].get(key)
        if not ent:
            continue
        opens = [b for b in ent["bills"] if b["status"] == "OPEN"]
        by_ccy: dict[str, float] = defaultdict(float)
        for b in opens:
            by_ccy[b["currency"]] += b["due_amount"]
        overdue = sum(1 for b in opens if b["due"] and b["due"] < today)
        counts = store.get("counts", {}).get(key, {})
        row = [SHORT(key), len(opens), *[counts.get(s, 0) for s in STATE_ORDER[:-1]], overdue, counts.get("PAID, UNRECONCILED", 0),
               ", ".join(f"{k} {v:,.2f}" for k, v in sorted(by_ccy.items())), ent.get("pulled", "")]
        for j, v in enumerate(row, start=1):
            c = ws.cell(r, j, v)
            if 2 <= j <= len(STATE_ORDER) + 2:
                c.alignment = Alignment(horizontal="center")
                if j >= 3 and isinstance(v, int) and v and STATE_ORDER[j - 3] in FILLS and j - 3 < len(STATE_ORDER) - 1:
                    c.fill = FILLS[STATE_ORDER[j - 3]]
        link = ws.cell(r, 1)
        link.hyperlink = f"#'{SHORT(key)}'!A1"
        link.font = Font(color="0563C1", underline="single", bold=True)
        tot["open"] += len(opens)
        tot["overdue"] += overdue
        for s in STATE_ORDER + ["PAID, UNRECONCILED"]:
            tot[s] += counts.get(s, 0)
        r += 1
    row = ["Total", tot["open"], *[tot[s] for s in STATE_ORDER[:-1]], tot["overdue"], tot["PAID, UNRECONCILED"], "", ""]
    for j, v in enumerate(row, start=1):
        c = ws.cell(r, j, v)
        c.font = Font(bold=True)
        if 2 <= j <= len(STATE_ORDER) + 2:
            c.alignment = Alignment(horizontal="center")

    r += 2
    ws.cell(r, 1, "What each status means").font = Font(bold=True)
    r += 1
    for s in STATE_ORDER:
        c = ws.cell(r, 1, s)
        c.fill = FILLS[s]
        c.font = Font(bold=True)
        ws.cell(r, 2, STATE_HELP[s])
        r += 1

    r += 1
    ws.cell(r, 1, "Manual actions: open bills paid from another entity's bank, and duplicates").font = Font(bold=True)
    r += 1
    any_ic = False
    for key in order:
        for b in store["entities"].get(key, {}).get("bills", []):
            if b.get("state") in ("INTERCOMPANY", "DUPLICATE?"):
                any_ic = True
                ws.cell(r, 1, SHORT(key))
                ws.cell(r, 2, f"{b['supplier']} · {b['date']} · {b['currency']} {b['due_amount']:,.2f} · {b['number']}: {b['why']}")
                r += 1
    if not any_ic:
        ws.cell(r, 2, "none")
    for col, w in zip("ABCDEFGHIJKLM", (12, 11, 14, 14, 11, 14, 11, 13, 12, 10, 12, 44, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"

    # ---- one tab per entity
    cols = ["Status", "Why / what to do", "Reconciled", "Bill date", "Due date", "Due status", "Supplier", "Invoice no",
            "Reference", "CCY", "Total", "Paid", "Credited", "Amount due", "Account(s)", "Line description",
            "Bank evidence", "Paid on", "Paid from", "Payment state", "Attachment", "Xero"]
    widths = [18, 58, 15, 11, 11, 13, 30, 20, 20, 6, 13, 11, 11, 13, 32, 50, 70, 11, 26, 40, 10, 8]
    REC_COL = 3
    for key in order:
        ent = store["entities"].get(key)
        if not ent:
            continue
        ws = wb.create_sheet(SHORT(key)[:31])
        ws["A1"] = f"{ents[key].title}: bills payable"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = (f"Xero pulled {ent.get('pulled', '')} UTC. Filtered to unreconciled: open bills oldest first, then "
                    f"paid bills whose payment is not reconciled. Clear the Reconciled filter for every bill paid since "
                    f"{ent.get('from', '')}.")
        ws["A2"].font = Font(size=9, italic=True, color="555555")
        hr = 4
        for j, h in enumerate(cols, start=1):
            c = ws.cell(hr, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
        ws.row_dimensions[hr].height = 30
        opens = sorted((b for b in ent["bills"] if b["status"] == "OPEN"), key=lambda b: (b["date"], b["supplier"]))
        paids = sorted((b for b in ent["bills"] if b["status"] == "PAID"), key=lambda b: (b["date"], b["supplier"]), reverse=True)
        r = hr + 1
        for b in opens + paids:
            if b["status"] == "OPEN":
                if not b["due"]:
                    due_state = ""
                elif b["due"] < today:
                    due_state = f"overdue {(date.fromisoformat(today) - date.fromisoformat(b['due'])).days}d"
                else:
                    due_state = "not yet due"
            else:
                due_state = ""
            reconciled = bool(b.get("reconciled"))
            row = [b.get("state", ""), b.get("why", ""), "reconciled" if reconciled else "not reconciled",
                   b["date"], b["due"], due_state, b["supplier"], b["number"],
                   b["reference"], b["currency"], b["total"], b["paid"], b["credited"], b["due_amount"],
                   "; ".join(b["accounts"]), b["description"], b.get("evidence", ""),
                   b.get("paid_on", ""), b.get("paid_from", ""), b.get("why", "") if b["status"] == "PAID" else "",
                   "yes" if b["attachment"] else "NO", "open"]
            if b["status"] == "PAID":
                row[1] = ""
            for j, v in enumerate(row, start=1):
                c = ws.cell(r, j, v)
                if j in (11, 12, 13, 14):
                    c.number_format = "#,##0.00"
                if j in (2, 16, 17, 20):
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                else:
                    c.alignment = Alignment(vertical="top")
            st = ws.cell(r, 1)
            st.fill = FILLS.get(b.get("state", ""), PatternFill())
            st.font = Font(bold=True)
            if due_state.startswith("overdue"):
                ws.cell(r, 6).font = Font(color="C00000", bold=True)
            if b["status"] == "PAID" and str(b.get("why", "")).startswith("PHANTOM"):
                ws.cell(r, 20).font = Font(color="C00000", bold=True)
            if not b["attachment"]:
                ws.cell(r, 21).font = Font(color="C00000")
            if reconciled:
                ws.row_dimensions[r].hidden = True   # what the Reconciled filter hides
            link = ws.cell(r, 22)
            if ent.get("short_code") and b["id"]:
                link.hyperlink = (f"https://go.xero.com/organisationlogin/default.aspx?shortcode={ent['short_code']}"
                                  f"&redirecturl=/AccountsPayable/View.aspx?InvoiceID={b['id']}")
                link.font = Font(color="0563C1", underline="single")
            else:
                link.value = ""
            r += 1
        ws.freeze_panes = ws.cell(hr + 1, 3)
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{max(r - 1, hr)}"
        # The filter as the reader opens it: unreconciled only. Excel keeps
        # the hidden rows in step with the criterion set here.
        ws.auto_filter.filterColumn.append(FilterColumn(colId=REC_COL - 1, filters=Filters(filter=["not reconciled"])))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------- CLI

def _publish(out: Path) -> None:
    from accounting_agent.publish import publish_to_drive
    print("Drive: " + (publish_to_drive(out, kind="bills") or "not published (no Drive folder configured)"))


def cmd_refresh(targets: list[str], date_from: str, publish: bool) -> int:
    import os
    from accounting_agent import readonly
    store = load_store()
    # Read-only against Xero by construction: the lock is on for the whole
    # pull. A lock a caller set stays set afterwards (and then Drive refuses
    # the publish, which is the point of a read-only session).
    lock_was_ours = not readonly.active()
    os.environ[readonly.ENV] = "1"
    try:
        for key in targets:
            print(f"{SHORT(key)}: pulling bills from Xero ...", flush=True)
            store["entities"][key] = pull_entity(key, date_from)
            ent = store["entities"][key]
            n_open = sum(1 for b in ent["bills"] if b["status"] == "OPEN")
            print(f"  {n_open} open, {len(ent['bills']) - n_open} paid since {date_from}, {len(ent['spends'])} spend money read")
        print("bank: reading the statements and the bank APIs ...", flush=True)
        assess(store)
    finally:
        if lock_was_ours:
            del os.environ[readonly.ENV]
    save_store(store)
    out = build_workbook(store)
    print(f"written {out.relative_to(ROOT)}")
    cmd_summary(store)
    if publish:
        _publish(out)
    return 0


def cmd_build(publish: bool) -> int:
    store = load_store()
    if not store["entities"]:
        sys.exit("no store yet; run refresh all first")
    assess(store)
    save_store(store)
    out = build_workbook(store)
    print(f"written {out.relative_to(ROOT)}")
    cmd_summary(store)
    if publish:
        _publish(out)
    return 0


def cmd_summary(store: dict | None = None) -> int:
    store = store or load_store()
    counts = store.get("counts", {})
    print(f"assessed {store.get('assessed', 'never')} · bank to {store.get('bank_upto', {})}")
    for key in ORDER():
        c = counts.get(key)
        if c is None:
            continue
        opens = sum(v for s, v in c.items() if s != "PAID")
        parts = ", ".join(f"{s.lower()} {c[s]}" for s in STATE_ORDER if c.get(s) and s != "PAID")
        print(f"{SHORT(key):10} {opens:4} open ({parts}) · {c.get('PAID', 0)} paid")
    return 0


def unpaid_message(store: dict) -> str:
    """The UNPAID bills, entity by entity, in the shape a payer needs:
    supplier · bill date · due date · due status · CCY amount · invoice number."""
    today = date.today()
    d = lambda iso: f"{int(iso[8:10])} {date.fromisoformat(iso).strftime('%b')}" if iso else ""   # noqa: E731
    parts = [f"unpaid bills, {today.day} {today.strftime('%b %Y')}"]
    total = 0
    for key in ORDER():
        rows = [b for b in store["entities"].get(key, {}).get("bills", []) if b.get("state") == "UNPAID"]
        if not rows:
            continue
        parts.append(f"\n*{SHORT(key).lower()}*")
        for b in sorted(rows, key=lambda b: (b["due"] or b["date"], b["supplier"])):
            if not b["due"]:
                due = ""
            elif b["due"] < today.isoformat():
                due = f" · overdue {(today - date.fromisoformat(b['due'])).days}d"
            else:
                due = " · not yet due"
            parts.append(f"• {b['supplier']} · {d(b['date'])} · due {d(b['due'])}{due} · "
                         f"{b['currency']} {b['due_amount']:,.2f} · {b['number']}".replace(" · due  ·", " ·"))
            total += 1
    if not total:
        parts.append("none")
    return "\n".join(parts)


def _recipient(who: str | None) -> str:
    """A Slack member ID: given, looked up by name in the [slack] tables, or
    [accounts_payable].unpaid_notify when nothing is given."""
    if not who:
        return str(config.accounts_payable().get("unpaid_notify") or "")
    if re.fullmatch(r"[UW][A-Z0-9]{6,}", who):
        return who
    needle = who.lower().strip()
    for uid, name in config.slack().allowed.items():
        full = (name or "").lower()
        if needle in (full, full.split(" ")[0], full.replace(" ", "-")):
            return uid
    return ""


def cmd_notify(who: str | None, dry_run: bool) -> int:
    import os
    uid = _recipient(who)
    if not uid:
        if not who:
            print("no recipient: [accounts_payable].unpaid_notify is blank, nothing sent")
            return 0
        names = ", ".join(config.slack().allowed.values()) or "(none configured)"
        sys.exit(f"unknown recipient {who!r}; a Slack member ID or one of: {names}")
    store = load_store()
    if not store.get("assessed"):
        sys.exit("no assessed store; run refresh all first")
    text = unpaid_message(store)
    if dry_run:
        print(f"DRY RUN, would DM {uid}:\n{text}")
        return 0
    token = os.environ.get("SLACK_BOT_TOKEN", "")
    if not token:
        sys.exit("SLACK_BOT_TOKEN not set (load .env)")
    from slack_sdk import WebClient
    WebClient(token=token).chat_postMessage(channel=uid, text=text)
    n = text.count("\n• ")
    print(f"sent unpaid list to {who or 'unpaid_notify'} ({uid}): {n} bills")
    return 0


def cmd_open(entity: str | None) -> int:
    store = load_store()
    want = config.entity(entity).slug if entity else None
    for key in ORDER():
        if want and key != want:
            continue
        ent = store["entities"].get(key)
        if not ent:
            continue
        print(f"== {SHORT(key)}")
        for b in sorted((b for b in ent["bills"] if b["status"] == "OPEN"), key=lambda b: (b.get("state", ""), b["date"])):
            print(f"  [{b.get('state', '?')}] {b['supplier']} · {b['date']} · {b['currency']} {b['due_amount']:,.2f} · {b['number']}")
            print(f"      {b.get('why', '')}")
            if b.get("evidence"):
                print(f"      bank: {b['evidence']}")
    return 0


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help", "help"):
        print(__doc__)
        return 0 if argv else 2
    cmd, rest = argv[0], argv[1:]
    publish = "--no-publish" not in rest
    date_from = rest[rest.index("--from") + 1] if "--from" in rest else None
    if cmd == "refresh" and rest:
        target = rest[0]
        if target.lower() == "all":
            targets = ORDER()
        else:
            try:
                targets = [config.entity(target).slug]
            except config.ConfigError as exc:
                sys.exit(f"{exc}, or all")
        store = load_store()
        prior = next((e.get("from") for e in store["entities"].values() if e.get("from")), None)
        return cmd_refresh(targets, date_from or prior or _records_from(), publish)
    if cmd == "build":
        return cmd_build(publish)
    if cmd == "summary":
        return cmd_summary()
    if cmd == "open":
        return cmd_open(rest[0] if rest and not rest[0].startswith("--") else None)
    if cmd == "notify":
        pos = [a for a in rest if not a.startswith("--")]
        return cmd_notify(pos[0] if pos else None, "--dry-run" in rest)
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
