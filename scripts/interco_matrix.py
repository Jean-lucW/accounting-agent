#!/usr/bin/env python
"""Intercompany summary matrix + one line-by-line tab per pair -> Excel.

The structure comes entirely from config/group.toml: the matrix axes are the
[[entities]] in file order, and every [[intercompany]] block gets its own
tab. Add an entity or a loan to the config and the next build has the new
row, column and tab; nothing here names an entity.

Sheet "summary matrix": entities on both axes. A cell is read
**column entity's books, its intercompany balance against the row entity**,
so (row=OpCo US, col=HoldCo) is what HoldCo's ledger says about OpCo US.
Every flavour configured for the two entities is summed, because the
question the matrix answers is whether the overall entity-to-entity
position agrees. The diagonal is blank. Below the matrix, one line per
configured pair (flavour by flavour) links to its tab.

Cell colour:
  GREEN  the two directions agree (A->B + B->A nets to ~0 in the reporting currency)
  GREY   no balance either way (both nil), or no account configured
  RED    the two directions disagree

Sign: debit-positive. A positive cell means the column entity is OWED by the
row entity; negative means it OWES. Agreement means the pair sums to zero.

Every configured pair gets its own tab, hyperlinked from the pair list and
(for the first flavour of each two entities) from the matrix cells: the
year's transactions, grouped by month (latest month first), with the two
ledgers side by side, each transaction against the one it matches in the
other entity, and 0 where the mirror is missing. A missing mirror is by
definition the cause of a difference. A group match (one journal against
several payments, several bills against one settlement) is one merged block
with the group's reporting-currency total in the Diff column. Column B is
the category: Interco, Reversal, FX Reval, Sweep JE (a consolidation or
sweep journal between interco accounts), CT (under the rounding tolerance),
and the tab opens filtered to Interco, the lines that have to reconcile with
the counterparty. A pair whose mirror account is missing (an empty account
code in the config) still gets a tab: a STRUCTURAL GAP banner and the
existing side's transactions, all one-sided.

Every build publishes the workbook to Drive (accounting_agent.publish, kind
"intercompany"): one copy under a fixed name that each build overwrites. A
read-only session cannot publish and says so; --no-publish skips it, which
is what deploy/run-interco.sh passes because it publishes itself outside the
read-only lock its Xero pull runs under.

Usage:
    .venv/bin/python scripts/interco_matrix.py --as-at 2026-08-31
    .venv/bin/python scripts/interco_matrix.py --as-at 2026-08-31 -o /tmp/interco.xlsx --no-publish
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import FilterColumn, Filters
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from accounting_agent import config  # noqa: E402
from interco_recon import (  # noqa: E402
    derive_rates, is_ct, is_fx_adjustment, is_sweep, match_sides, month_end, pull_balances, pull_txns, rate,
    reporting, settings, to_rep, wrong_flavour_hints,
)

WORKBOOK = "Intercompany Reconciliation.xlsx"
DEFAULT_OUT = ROOT / "data" / "reports" / WORKBOOK
MATRIX_SHEET = "summary matrix"

GREEN = PatternFill("solid", fgColor="C6E0B4")
RED = PatternFill("solid", fgColor="F4B6B6")
GREY = PatternFill("solid", fgColor="D9D9D9")
HEAD = PatternFill("solid", fgColor="203864")
SUBHEAD = PatternFill("solid", fgColor="D6DCE4")
MONTHFILL = PatternFill("solid", fgColor="EDF2F9")
GAPFILL = PatternFill("solid", fgColor="FCE4D6")
A_FILL = PatternFill("solid", fgColor="F2F7F2")
B_FILL = PatternFill("solid", fgColor="F7F2F2")
THIN = Side(style="thin", color="9AA0A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00;[Red]-#,##0.00'
LINKFONT = Font(color="0563C1", underline="single")


# ------------------------------------------------------------ group layout

def order() -> list[str]:
    """Matrix axis order: the entities in config file order."""
    return [e.key for e in config.entities()]


def label(key: str) -> str:
    e = config.entity(key)
    return f"{e.short} ({e.key})" if e.short != e.key else e.key


def short(key: str) -> str:
    return config.entity(key).short


def base(key: str) -> str:
    return config.entity(key).base_currency


def agrees(rep_break: float, a: float, b: float, cross: bool) -> bool:
    """Same-base pairs must agree to the rounding tolerance. Cross-base pairs
    are compared through an estimated cross rate, so a small relative band
    absorbs the rate estimate; without it every FX pair shows red for pure
    rounding. Both tolerances come from [intercompany_settings]."""
    s = settings()
    tol_abs, tol_rel = float(s["tolerance_abs"]), float(s["tolerance_rel"])
    if not cross:
        return abs(rep_break) <= tol_abs
    return abs(rep_break) <= max(tol_abs, tol_rel * max(abs(a), abs(b)))


def link_to_sheet(cell, sheet: str) -> None:
    """An in-workbook hyperlink. openpyxl treats a plain "#'Sheet'!A1" string
    as an external target, which Excel and Google Sheets cannot follow; the
    `location` form is the internal link both open. Google Sheets rewrites
    the link as =HYPERLINK(url, display) on import and, with no display text,
    shows the url instead of the figure, so the display is the cell's own
    value, formatted as the cell would show it. Set the value first."""
    v = cell.value
    display = f"{v:,.2f}" if isinstance(v, (int, float)) else ("" if v is None else str(v))
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=display)


_BAD = re.compile(r"[\[\]:*?/\\]")


def sheet_names(pairs) -> dict:
    """{pair.key: sheet name}, one per configured pair. Excel sheet names are at
    most 31 characters, may not contain []:*?/\\ , may not start or end with
    an apostrophe, and are unique case-insensitively. The name is built from
    the two entities' short names and the flavour; a collision after
    truncation gets a numeric suffix."""
    used = {MATRIX_SHEET.lower()}
    out = {}
    for p in pairs:
        raw = _BAD.sub("-", f"{short(p.a)}-{short(p.b)} {p.flavour}").strip().strip("'") or "pair"
        name = raw[:31].rstrip()
        n = 2
        while name.lower() in used:
            sfx = f" ({n})"
            name = raw[:31 - len(sfx)].rstrip() + sfx
            n += 1
        used.add(name.lower())
        out[p.key] = name
    return out


# ------------------------------------------------------------ data

def collect(as_at: str, year_from: str, only: list[str] | None = None):
    """Balances (monthly) and full-year transactions for every interco account.

    `only` limits the entities queried: useful when one entity has burned its
    Xero daily quota (5,000 calls/entity/day); the rest still reconcile.
    """
    keys = order()
    codes: dict[str, set[str]] = {k: set() for k in keys}
    for p in config.intercompany_pairs():
        for ent in (p.a, p.b):
            if p.account(ent):
                codes[ent].add(p.account(ent))
    ents = [e for e in keys if (not only or e in only) and codes[e]]

    # enough monthly columns to reach back to the opening balance
    a, y = datetime.date.fromisoformat(as_at), datetime.date.fromisoformat(year_from)
    months = max(2, (a.year - y.year) * 12 + a.month - y.month + 2)

    # Two sequential phases, 5-wide each. Running both together bursts past
    # Xero's 60 calls/min per entity and 429s: the transaction rebuild alone
    # is ~6 endpoints per entity.
    with ThreadPoolExecutor(5) as ex:
        bals = dict(zip(ents, ex.map(lambda e: pull_balances(e, as_at, months), ents)))
    if ents:
        time.sleep(20)
    with ThreadPoolExecutor(5) as ex:
        # to the month end of as_at: the Balance Sheet column is the month end
        # and already holds anything dated forward to it
        txns = dict(zip(ents, ex.map(lambda e: pull_txns(e, codes[e], year_from, month_end(as_at)), ents)))
    for e in keys:                       # entities not queried look empty, not missing
        bals.setdefault(e, {})
        txns.setdefault(e, [])

    rates = derive_rates(txns, as_at)
    labels = next((list(v) for e in ents for v in bals[e].values() if v), [])
    closing = labels[0] if labels else None
    # The opening balance is the month end IMMEDIATELY BEFORE the transaction
    # window, not the oldest column: `months` carries a margin, so labels[-1]
    # overshoots by a month and the tab's Unexplained line then shows that
    # month's movement, above all on flavours that are swept to nil at a
    # period end and so differ wildly one month either side of it. Match by
    # date, fall back to the oldest column.
    opening = labels[-1] if labels else None
    if labels:
        prev_end = datetime.date.fromisoformat(year_from) - datetime.timedelta(days=1)
        want = f"{prev_end.day} {prev_end.strftime('%b %Y')}"
        opening = next((lb for lb in labels if lb == want), opening)
    return bals, txns, rates, closing, opening


# Column B. The tab opens filtered to Interco (plus the blank month rows), so
# what is left on screen is the business with the counterparty that has to
# reconcile line by line; the rest is set aside by kind, not deleted.
CAT_INTERCO, CAT_REVERSAL, CAT_FX, CAT_SWEEP, CAT_CT = "Interco", "Reversal", "FX Reval", "Sweep JE", "CT"


def category(items, rates, reversed_=False, fx=False):
    if reversed_:
        return CAT_REVERSAL
    if fx:
        return CAT_FX
    if items and all(is_sweep(t) for t in items):
        return CAT_SWEEP
    if items and all(is_ct(t, rates) for t in items):
        return CAT_CT
    return CAT_INTERCO


def pair_rows(a_ent, b_ent, res, rates, hints=None):
    """One display block per economic event, both ledgers side by side.

    `res` is match_sides() for the pair. A block is (date, a_items, b_items,
    note, category). Most blocks are one item a side; a group or probable
    match carries several on one or both sides and is drawn as a merged
    block. Every transaction appears exactly once, so the columns total to
    each side's movement for the year.
    """
    hints = hints or {}
    rep = reporting()

    def tot(ts):
        return sum(to_rep(t, rates) for t in ts)

    out = []
    for a, b, why in res["matched"]:
        note = "" if why == "exact" else ("matched on base ccy" if why == "base" else "matched cross-currency")
        if a.date != b.date:
            note = (note + "; " if note else "") + "TIMING: dates differ"
        out.append((min(a.date, b.date), [a], [b], note, category([a, b], rates)))
    for ga, gb, why in res["groups"]:
        d0 = min(t.date for t in ga + gb)
        out.append((d0, sorted(ga, key=lambda t: t.date), sorted(gb, key=lambda t: t.date),
                    f"GROUP MATCH {len(ga)} in {short(a_ent)} against {len(gb)} in {short(b_ent)}",
                    category(ga + gb, rates)))
    for ga, gb, why in res["probable"]:
        d0 = min(t.date for t in ga + gb)
        out.append((d0, sorted(ga, key=lambda t: t.date), sorted(gb, key=lambda t: t.date),
                    f"PROBABLE MATCH - same supplier or subject, amounts differ by {rep} "
                    f"{tot(ga) + tot(gb):,.2f}, investigate", CAT_INTERCO))
    for t in res["unmatched_a"]:
        out.append((t.date, [t], [], hints.get(id(t)) or f"ONE-SIDED - no entry in {short(b_ent)}",
                    category([t], rates)))
    for t in res["unmatched_b"]:
        out.append((t.date, [], [t], hints.get(id(t)) or f"ONE-SIDED - no entry in {short(a_ent)}",
                    category([t], rates)))
    for t in res["fx_adjustments"]["a"]:
        out.append((t.date, [t], [], "FX revaluation (one-sided by design)", CAT_FX))
    for t in res["fx_adjustments"]["b"]:
        out.append((t.date, [], [t], "FX revaluation (one-sided by design)", CAT_FX))
    for side, pairs in (("a", res["contras"]["a"]), ("b", res["contras"]["b"])):
        for x, y in pairs:
            for t in (x, y):
                out.append((t.date, [t] if side == "a" else [], [t] if side == "b" else [],
                            "reversed - posting and its reversal, nets to nil", CAT_REVERSAL))
    return out


def gap_rows(a_ent, b_ent, missing, a_rows, b_rows, rates):
    """A pair with no mirror account: every existing-side line is one-sided
    by construction (break class G)."""
    out = []
    for t in a_rows:
        out.append((t.date, [t], [], f"STRUCTURAL GAP - no mirror account in {short(missing)}",
                    CAT_FX if is_fx_adjustment(t) else category([t], rates)))
    for t in b_rows:
        out.append((t.date, [], [t], f"STRUCTURAL GAP - no mirror account in {short(missing)}",
                    CAT_FX if is_fx_adjustment(t) else category([t], rates)))
    return out


# Column layout of a pair tab. B is the Category the tab is filtered on.
COL = {"date": 1, "cat": 2, "a_ac": 3, "a_src": 4, "a_desc": 5, "a_amt": 6, "a_rep": 7,
       "b_ac": 8, "b_src": 9, "b_desc": 10, "b_amt": 11, "b_rep": 12, "diff": 13, "note": 14}
NCOL = 14


def _spans(k: int, R: int):
    """Row offsets (start, end) for k items laid across R rows, evenly."""
    return [(i * R // k, (i + 1) * R // k - 1) for i in range(k)]


def _put(ws, r0, r1, col, value, fmt=None, font=None, fill=None):
    c = ws.cell(r0, col, value)
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    if fill:
        for r in range(r0, r1 + 1):
            ws.cell(r, col).fill = fill
    if r1 > r0:
        ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
        c.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left")
    return c


def write_pair_sheet(wb, sheet, pair, rows, rates, opening_a, opening_b, closing_a, closing_b, as_at,
                     missing: str | None = None):
    a_ent, b_ent = pair.a, pair.b
    ws = wb.create_sheet(sheet)
    ba, bb = base(a_ent), base(b_ent)
    rep = reporting()
    tol = float(settings()["tolerance_abs"])
    code_a = pair.account(a_ent) or "none"
    code_b = pair.account(b_ent) or "none"

    ws["A1"] = f"{label(a_ent)}   <->   {label(b_ent)}   ·   {pair.label} ({pair.flavour})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Transactions to {month_end(as_at)} (the month end the balances are struck at), grouped by month, "
                f"latest first. Each row pairs a transaction with its mirror in the other ledger; 0 means the mirror "
                f"is missing. A merged block is a group match: several lines one side against one or several the other.")
    ws["A3"] = (f"Debit-positive. A matched pair or group should net to zero in {rep}; the Diff column is the residual. "
                "Column B is the category: Interco (business with the counterparty, must reconcile line by line), "
                "Reversal (a posting undone by its own reversal), FX Reval (one-sided by design), "
                "Sweep JE (a consolidation or sweep journal between interco accounts), CT (under the rounding "
                "tolerance). The tab opens filtered to Interco.")
    for r in (2, 3):
        ws[f"A{r}"].font = Font(size=9, italic=True, color="555555")
    if missing:
        g = ws["A4"]
        g.value = (f"STRUCTURAL GAP: {config.entity(missing).xero_name} has no mirror account for this pair "
                   f"(empty account code in config/group.toml [[intercompany]]). Open the account in Xero, add its "
                   f"code to the config, and this tab reconciles on the next build.")
        g.font = Font(bold=True, color="C00000")
        for j in range(1, NCOL + 1):
            ws.cell(4, j).fill = GAPFILL

    hdr_top = 5
    ws.merge_cells(start_row=hdr_top, start_column=COL["a_ac"], end_row=hdr_top, end_column=COL["a_rep"])
    ws.merge_cells(start_row=hdr_top, start_column=COL["b_ac"], end_row=hdr_top, end_column=COL["b_rep"])
    ca = ws.cell(hdr_top, COL["a_ac"], f"{label(a_ent)}  [{code_a}]  ({ba})")
    cb = ws.cell(hdr_top, COL["b_ac"], f"{label(b_ent)}  [{code_b}]  ({bb})")
    for c in (ca, cb):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center")

    hdr = ["Date", "Category", "A/c", "Source", "Description", f"Amount {ba}", rep,
           "A/c", "Source", "Description", f"Amount {bb}", rep, f"Diff {rep}", "Note"]
    row = hdr_top + 1
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row, j, h)
        c.font = Font(bold=True)
        c.fill = SUBHEAD
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center")

    by_month = defaultdict(list)
    for blk in rows:
        by_month[blk[0][:7]].append(blk)

    r = row + 1
    tot_a = tot_b = 0.0
    hidden_rows = []
    for month in sorted(by_month, reverse=True):
        mc = ws.cell(r, 1, datetime.date.fromisoformat(month + "-01").strftime("%B %Y"))
        mc.font = Font(bold=True)
        for j in range(1, NCOL + 1):
            ws.cell(r, j).fill = MONTHFILL
        r += 1
        ma = mb = 0.0
        for d, A, B, note, cat in sorted(by_month[month], key=lambda x: x[0], reverse=True):
            R = max(len(A), len(B), 1)
            r1 = r + R - 1
            rev = cat == CAT_REVERSAL
            _put(ws, r, r1, COL["date"], d)
            _put(ws, r, r1, COL["cat"], cat,
                 font=Font(bold=True, color="C00000") if rev else Font(color="000000" if cat == CAT_INTERCO else "7F7F7F"))
            ua_tot = ub_tot = 0.0
            for items, keys, fill, base_ccy in ((A, ("a_ac", "a_src", "a_desc", "a_amt", "a_rep"), A_FILL, ba),
                                                 (B, ("b_ac", "b_src", "b_desc", "b_amt", "b_rep"), B_FILL, bb)):
                if not items:
                    for k in keys:
                        _put(ws, r, r1, COL[k], 0 if k.endswith(("amt", "rep")) else None,
                             fmt=MONEY if k.endswith(("amt", "rep")) else None, fill=fill)
                    continue
                for (o0, o1), t in zip(_spans(len(items), R), items):
                    in_rep = t.base * rate(rates, base_ccy)
                    _put(ws, r + o0, r + o1, COL[keys[0]], t.code, fill=fill)
                    _put(ws, r + o0, r + o1, COL[keys[1]], t.source, fill=fill)
                    _put(ws, r + o0, r + o1, COL[keys[2]],
                         (t.party + " - " + t.desc)[:70] if t.party else t.desc[:70], fill=fill)
                    _put(ws, r + o0, r + o1, COL[keys[3]], round(t.base, 2), fmt=MONEY, fill=fill)
                    _put(ws, r + o0, r + o1, COL[keys[4]], round(in_rep, 2), fmt=MONEY, fill=fill)
                    if items is A:
                        ua_tot += in_rep
                        ma += t.base
                    else:
                        ub_tot += in_rep
                        mb += t.base
            dc = _put(ws, r, r1, COL["diff"], round(ua_tot + ub_tot, 2), fmt=MONEY)
            if abs(ua_tot + ub_tot) > tol and cat == CAT_INTERCO:
                dc.font = Font(bold=True, color="C00000")
            nc = _put(ws, r, r1, COL["note"], note)
            nc.font = Font(size=9, color="806000" if note else "000000")
            if cat != CAT_INTERCO:
                hidden_rows.extend(range(r, r1 + 1))
            r = r1 + 1
        sc = ws.cell(r, COL["a_desc"], "month movement")
        sc.font = Font(bold=True, italic=True)
        ws.cell(r, COL["a_amt"], round(ma, 2)).number_format = MONEY
        ws.cell(r, COL["b_amt"], round(mb, 2)).number_format = MONEY
        for j in (COL["a_amt"], COL["b_amt"]):
            ws.cell(r, j).font = Font(bold=True, italic=True)
        tot_a += ma
        tot_b += mb
        r += 2
    if not rows:
        ws.cell(r, 1, "no transactions in the period").font = Font(italic=True, color="7F7F7F")
        r += 2
    last_data = r - 2

    # Filter on Category: Interco plus blanks (the month header and movement
    # rows). Excel and Sheets read the criterion, and the other rows are
    # hidden as well so the file opens filtered.
    if last_data > row and rows:
        ws.auto_filter.ref = f"A{row}:{get_column_letter(NCOL)}{last_data}"
        ws.auto_filter.filterColumn.append(
            FilterColumn(colId=COL["cat"] - 1, filters=Filters(blank=True, filter=[CAT_INTERCO])))
        for hr in hidden_rows:
            ws.row_dimensions[hr].hidden = True

    r += 1
    ws.cell(r, COL["a_desc"], "Opening balance").font = Font(bold=True)
    ws.cell(r, COL["a_amt"], round(opening_a, 2)).number_format = MONEY
    ws.cell(r, COL["b_amt"], round(opening_b, 2)).number_format = MONEY
    ws.cell(r + 1, COL["a_desc"], "Movement in period").font = Font(bold=True)
    ws.cell(r + 1, COL["a_amt"], round(tot_a, 2)).number_format = MONEY
    ws.cell(r + 1, COL["b_amt"], round(tot_b, 2)).number_format = MONEY
    ws.cell(r + 2, COL["a_desc"], "Closing per Xero").font = Font(bold=True)
    ws.cell(r + 2, COL["a_amt"], round(closing_a, 2)).number_format = MONEY
    ws.cell(r + 2, COL["b_amt"], round(closing_b, 2)).number_format = MONEY
    ws.cell(r + 3, COL["a_desc"], "Unexplained (should be 0)").font = Font(bold=True)
    for col, op, mv, cl in ((COL["a_amt"], opening_a, tot_a, closing_a), (COL["b_amt"], opening_b, tot_b, closing_b)):
        gap = round(cl - op - mv, 2)
        c = ws.cell(r + 3, col, gap)
        c.number_format = MONEY
        c.font = Font(bold=True, color="C00000" if abs(gap) > 0.01 else "375623")
    ws.cell(r + 5, COL["a_desc"], f"Closing break ({rep})").font = Font(bold=True)
    brk = ws.cell(r + 5, COL["a_amt"],
                  "n/a (no mirror account)" if missing
                  else round(closing_a * rate(rates, ba) + closing_b * rate(rates, bb), 2))
    brk.number_format = MONEY
    brk.font = Font(bold=True, color="C00000")

    for j, w in enumerate((11, 10, 7, 13, 40, 15, 14, 7, 13, 40, 15, 14, 13, 40), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)

    back = ws.cell(1, NCOL, "<- back to summary matrix")
    link_to_sheet(back, MATRIX_SHEET)
    back.font = LINKFONT
    return ws.title


# ------------------------------------------------------------ workbook

def build_workbook(bals, txns, rates, closing, opening, as_at: str, out_path) -> dict:
    """Write the workbook from already-collected data. Pure: no API calls, so
    it can be exercised offline. Returns {pair name: sheet name}."""
    ORDER = order()
    rep = reporting()
    pairs = config.intercompany_pairs()
    names = sheet_names(pairs)
    wb = Workbook()
    ws = wb.active
    ws.title = MATRIX_SHEET

    def bal(ent, code, col):
        return (bals.get(ent, {}).get(code) or {}).get(col, 0.0) if code else 0.0

    positions: dict[tuple[str, str], float] = {}
    for p in pairs:
        for owner, other in ((p.a, p.b), (p.b, p.a)):
            code = p.account(owner)
            if code:
                positions[(owner, other)] = positions.get((owner, other), 0.0) + bal(owner, code, closing)

    # ------------------------------------------------- match every pair first
    # (the wrong-flavour hints need every flavour of two entities at once)
    results, rows_of = {}, {}
    for p in pairs:
        a_code, b_code = p.account(p.a), p.account(p.b)
        a_rows = [t for t in txns.get(p.a, []) if a_code and t.code == a_code]
        b_rows = [t for t in txns.get(p.b, []) if b_code and t.code == b_code]
        if p.complete:
            results[p.key] = match_sides(a_rows, b_rows, rates)
        else:
            rows_of[p.key] = (a_rows, b_rows)
    hints = wrong_flavour_hints([(p, results[p.key]) for p in pairs if p.key in results], rates)

    # ------------------------------------------------- one tab per configured pair
    tabs: dict = {}
    summary = []
    for p in pairs:
        a_code, b_code = p.account(p.a), p.account(p.b)
        ca, cb = bal(p.a, a_code, closing), bal(p.b, b_code, closing)
        oa, ob = bal(p.a, a_code, opening), bal(p.b, b_code, opening)
        if p.complete:
            rows = pair_rows(p.a, p.b, results[p.key], rates, hints)
            missing = None
        else:
            missing = p.a if not a_code else p.b
            rows = gap_rows(p.a, p.b, missing, *rows_of[p.key], rates)
        tabs[p.key] = write_pair_sheet(wb, names[p.key], p, rows, rates, oa, ob, ca, cb, as_at, missing)
        summary.append((p, ca, cb, missing))

    # ------------------------------------------------------------ summary matrix
    ws["A1"] = f"{config.company_name()} intercompany summary matrix"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"As at {as_at}  (Xero balance column '{closing}')"
    ws["A3"] = ("Cell = the COLUMN entity's books: its intercompany balance against the ROW entity, "
                "all flavours summed, debit-positive in the column entity's base currency. "
                "Click a cell or a pair below for the transaction detail.")
    others = ", ".join(f"{k}={v:.5g}" for k, v in sorted(rates.items()) if k != rep)
    ws["A4"] = ("Green = the two directions agree   |   Red = they disagree   |   Grey = nil, or no account.   "
                f"Cross-currency breaks are translated to {rep} at {others or 'no cross rates'} "
                "(median of Xero document rates within 45 days of the balance date) and move with that rate; "
                "the authoritative figure is the per-transaction detail on the pair tab.")
    for r in (2, 3, 4):
        ws[f"A{r}"].font = Font(size=9, italic=True, color="555555")

    top = 6
    hc = ws.cell(top, 1, "ROW entity  →  as recorded by COLUMN entity")
    hc.font = Font(bold=True, size=9)
    hc.fill = SUBHEAD
    hc.alignment = Alignment(wrap_text=True, vertical="center")
    for j, col_ent in enumerate(ORDER, start=2):
        c = ws.cell(top, j, f"{label(col_ent)}\n({base(col_ent)})")
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[top].height = 32

    first_tab: dict[frozenset, str] = {}
    for p in pairs:
        first_tab.setdefault(frozenset((p.a, p.b)), tabs[p.key])

    for i, row_ent in enumerate(ORDER, start=top + 1):
        rc = ws.cell(i, 1, label(row_ent))
        rc.font = Font(bold=True, color="FFFFFF")
        rc.fill = HEAD
        rc.border = BOX
        for j, col_ent in enumerate(ORDER, start=2):
            cell = ws.cell(i, j)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="right")
            if row_ent == col_ent:
                cell.fill = PatternFill("solid", fgColor="404040")
                continue
            here, there = positions.get((col_ent, row_ent)), positions.get((row_ent, col_ent))
            tab = first_tab.get(frozenset((row_ent, col_ent)))
            if here is None and there is None:
                cell.value = "no a/c"
                cell.fill = GREY
                cell.font = Font(size=9, italic=True, color="666666")
                if tab:                      # configured, but no account on either side yet
                    link_to_sheet(cell, tab)
                continue
            if here is None:
                cell.value = "no mirror a/c"
                cell.fill = RED
                cell.font = Font(size=9, bold=True, color="C00000")
                if tab:
                    link_to_sheet(cell, tab)
                continue
            cell.value = round(here, 2)
            cell.number_format = MONEY
            ua = here * rate(rates, base(col_ent))
            ub = (there or 0.0) * rate(rates, base(row_ent))
            if abs(here) < 0.005 and abs(there or 0.0) < 0.005:
                cell.fill = GREY
            elif there is not None and agrees(ua + ub, ua, ub, base(col_ent) != base(row_ent)):
                cell.fill = GREEN
            else:
                cell.fill = RED
            if tab:
                link_to_sheet(cell, tab)
                cell.font = Font(color="0563C1", underline="single", bold=True)

    ws.column_dimensions["A"].width = 34
    for j in range(2, max(len(ORDER), 8) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 19
    ws.freeze_panes = ws.cell(top + 1, 2)

    # ------------------------------------------------- every configured pair
    r = top + len(ORDER) + 2
    ws.cell(r, 1, "Pairs (one tab each)").font = Font(bold=True)
    r += 1
    for j, h in enumerate(["Pair / tab", "Flavour", "A entity", "A a/c", "A balance", "B entity", "B a/c",
                           "B balance", f"Break {rep}", "Status"], start=1):
        c = ws.cell(r, j, h)
        c.font = Font(bold=True)
        c.fill = SUBHEAD
        c.border = BOX
    r += 1
    any_break = False
    for p, ca, cb, missing in summary:
        c = ws.cell(r, 1, tabs[p.key])
        link_to_sheet(c, tabs[p.key])
        c.font = LINKFONT
        ws.cell(r, 2, p.flavour)
        ws.cell(r, 3, p.a)
        ws.cell(r, 4, p.account(p.a) or "none")
        ws.cell(r, 5, round(ca, 2)).number_format = MONEY
        ws.cell(r, 6, p.b)
        ws.cell(r, 7, p.account(p.b) or "none")
        ws.cell(r, 8, round(cb, 2)).number_format = MONEY
        if missing:
            ws.cell(r, 9, "n/a")
            st = ws.cell(r, 10, f"STRUCTURAL GAP: no mirror account in {missing}")
            st.font = Font(bold=True, color="C00000")
            for j in range(1, 11):
                ws.cell(r, j).fill = GAPFILL
            any_break = True
        else:
            ua, ub = ca * rate(rates, base(p.a)), cb * rate(rates, base(p.b))
            ws.cell(r, 9, round(ua + ub, 2)).number_format = MONEY
            if abs(ca) < 0.005 and abs(cb) < 0.005:
                ws.cell(r, 10, "nil both sides")
            elif agrees(ua + ub, ua, ub, base(p.a) != base(p.b)):
                ws.cell(r, 10, "agreed").fill = GREEN
            else:
                n_un = len(results[p.key]["unmatched_a"]) + len(results[p.key]["unmatched_b"])
                st = ws.cell(r, 10, f"BREAK ({n_un} one-sided)")
                st.fill = RED
                st.font = Font(bold=True, color="C00000")
                any_break = True
        r += 1
    r += 1
    if not any_break:
        ws.cell(r, 1, "none - every pair agrees").font = Font(italic=True, color="375623")

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {p.name: tabs[p.key] for p in pairs}


def build(as_at: str, out_path, year_from: str, only: list[str] | None = None) -> str:
    bals, txns, rates, closing, opening = collect(as_at, year_from, only)
    build_workbook(bals, txns, rates, closing, opening, as_at, out_path)
    return str(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--as-at", default=datetime.date.today().isoformat())
    ap.add_argument("--from", dest="year_from", default=None,
                    help="first transaction date for the pair tabs (default: first day of the "
                         "financial year containing --as-at, [company] financial_year_end)")
    ap.add_argument("-o", "--out", default=None, help=f"default data/reports/{WORKBOOK}")
    ap.add_argument("--entities", default=None,
                    help="comma-separated subset of entity keys (default: all in config). "
                         "Use when an entity has exhausted its Xero daily call quota.")
    ap.add_argument("--no-publish", action="store_true",
                    help="build only; do not push the workbook to the intercompany folder on Drive")
    args = ap.parse_args()
    only = [config.entity(x.strip()).key for x in args.entities.split(",")] if args.entities else None
    year_from = args.year_from or config.fy_start_for(
        datetime.date.fromisoformat(args.as_at)).isoformat()
    out = args.out or str(DEFAULT_OUT)
    print(f"wrote {build(args.as_at, out, year_from, only)}")
    if args.no_publish:
        print("Drive: NOT published (--no-publish)")
    elif only:
        # a subset built around a rate-capped entity is a partial picture; the
        # Drive copy is always the full matrix
        print("Drive: NOT published, --entities builds are partial and stay on the server")
    else:
        from accounting_agent.publish import publish_to_drive
        print("Drive: " + (publish_to_drive(out, kind="intercompany") or "not published"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
